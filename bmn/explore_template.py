# %% [markdown]
# # Exploration du template Excel
#
# Liste, onglet par onglet : cellules à formule (à NE PAS toucher), cellules
# d'input non vides (valeurs d'exemple / libellés) et cellules vides dans la
# zone utilisée (candidates au remplissage). Sert à construire le mapping
# dans `config.py`.
#
# Usage : `python explore_template.py [chemin.xlsx] [--sheet NOM] [--export rapport.csv]`

# %%
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

import config
from daily_report import is_formula_cell


def scan_sheet(ws: Worksheet) -> pd.DataFrame:
    """Une ligne par cellule de la zone utilisée : coord, kind, value."""
    rows = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            if isinstance(cell, MergedCell):
                kind = "merged"
            elif is_formula_cell(cell):
                kind = "formula"
            elif cell.value is None:
                kind = "empty"
            else:
                kind = "input"
            # Une MergedCell n'a pas de column_letter : on le calcule depuis l'index.
            col = get_column_letter(cell.column)
            rows.append({
                "sheet": ws.title, "coord": cell.coordinate,
                "row": cell.row, "col": col,
                "kind": kind,
                "value": str(cell.value) if cell.value is not None else "",
                "number_format": cell.number_format,
            })
    return pd.DataFrame(rows)


def formula_columns_summary(df: pd.DataFrame) -> pd.DataFrame:
    """Colonnes contenant des formules, avec plage de lignes : repère les colonnes calculées."""
    f = df[df.kind == "formula"]
    if f.empty:
        return f
    return (f.groupby(["sheet", "col"])
             .agg(first_row=("row", "min"), last_row=("row", "max"), n=("row", "size"),
                  example=("value", "first"))
             .reset_index())


def describe_workbook(path: Path, sheet: str | None = None) -> pd.DataFrame:
    wb = load_workbook(path, data_only=False)
    print(f"Fichier : {path}")
    print(f"Onglets : {wb.sheetnames}")
    if wb.defined_names:
        print(f"Noms définis : {list(wb.defined_names.keys())}")

    sheets = [wb[sheet]] if sheet else wb.worksheets
    frames = []
    for ws in sheets:
        df = scan_sheet(ws)
        frames.append(df)
        counts = df.kind.value_counts().to_dict()
        print(f"\n=== Onglet '{ws.title}' | dimensions {ws.dimensions} | {counts}")
        if ws.merged_cells.ranges:
            plages = sorted(str(r) for r in ws.merged_cells.ranges)
            apercu = ", ".join(plages[:10])
            reste = f" … (+{len(plages) - 10})" if len(plages) > 10 else ""
            print(f"Plages fusionnées : {len(plages)} — {apercu}{reste}")

        fc = formula_columns_summary(df)
        bulk_cols = set(fc[fc.n > 5].col) if not fc.empty else set()
        print("\n-- Cellules à FORMULE isolées (protégées) --")
        for _, r in df[(df.kind == "formula") & (~df.col.isin(bulk_cols))].iterrows():
            print(f"  {r.coord:>6}  {r.value}")
        if bulk_cols:
            print("\n-- Colonnes calculées (plus de 5 formules, protégées) --")
            print(fc[fc.col.isin(bulk_cols)].to_string(index=False))

        print("\n-- Cellules d'INPUT non vides (libellés / exemples) --")
        for _, r in df[df.kind == "input"].iterrows():
            print(f"  {r.coord:>6}  {r.value!s:.60}")

        empties = df[df.kind == "empty"]
        print(f"\n-- Cellules VIDES dans la zone utilisée : {len(empties)} "
              f"(ex. {', '.join(empties.coord.head(15))}{'...' if len(empties) > 15 else ''})")
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="Explore un template Excel (formules vs inputs).")
    p.add_argument("path", nargs="?", type=Path, default=config.TEMPLATE_PATH)
    p.add_argument("--sheet", default=None, help="Limiter à un onglet.")
    p.add_argument("--export", type=Path, default=None,
                   help="Exporter l'inventaire complet des cellules en CSV.")
    a = p.parse_args(argv)
    if not a.path.exists():
        print(f"Template introuvable : {a.path}", file=sys.stderr)
        return 1
    df = describe_workbook(a.path, a.sheet)
    if a.export:
        df.to_csv(a.export, index=False, sep=";", encoding="utf-8-sig")
        print(f"\nInventaire exporté : {a.export}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
