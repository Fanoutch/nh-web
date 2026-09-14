# %% [markdown]
# # Fixtures factices
#
# Génère un template Excel AVEC formules et un CSV bidon dont les colonnes
# correspondent aux placeholders de `config.py`. Permet de tester le pipeline
# de bout en bout en attendant le vrai template / CSV.
#
# Usage : `python make_fixtures.py [--template] [--csv] [--date 2026-09-03]`
# (sans option : les deux). Le template n'est PAS écrasé s'il existe déjà,
# sauf avec `--force` : ne jamais écraser le vrai template par erreur.

# %%
from __future__ import annotations

import argparse
import sys
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

import config

FAKE_COLUMNS = [
    "COLONNE_A_DEFINIR_1",   # date rapport
    "COLONNE_A_DEFINIR_2",   # site
    "COLONNE_A_DEFINIR_3",   # référence
    "COLONNE_A_DEFINIR_4",   # libellé
    "COLONNE_A_DEFINIR_5",   # quantité
    "COLONNE_A_DEFINIR_6",   # prix unitaire
    "COLONNE_IGNOREE",       # colonne présente dans le CSV mais non mappée
]


def make_template(path: Path, sheet: str = "Saisie", table_start: int = 7,
                  table_rows: int = 50) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    bold = Font(bold=True)

    # En-tête : cellules d'input B2..B4 + formule en B5
    ws["A1"] = "Rapport quotidien"; ws["A1"].font = bold
    ws["A2"] = "Date"; ws["B2"].number_format = "DD/MM/YYYY"
    ws["A3"] = "Site"
    ws["A4"] = "Total quantité (import)"
    ws["A5"] = "Total montant (formule)"
    ws["B5"] = f"=SUM(E{table_start}:E{table_start + table_rows - 1})"

    # Tableau : A..D inputs, E formule, ligne de total sous le tableau
    headers = ["Référence", "Libellé", "Quantité", "PU", "Montant"]
    for i, h in enumerate(headers):
        c = ws.cell(row=table_start - 1, column=i + 1, value=h)
        c.font = bold
    for r in range(table_start, table_start + table_rows):
        ws[f"E{r}"] = f'=IF(C{r}="","",C{r}*D{r})'
    total_row = table_start + table_rows
    ws[f"A{total_row}"] = "TOTAL"; ws[f"A{total_row}"].font = bold
    ws[f"C{total_row}"] = f"=SUM(C{table_start}:C{total_row - 1})"
    ws[f"E{total_row}"] = f"=SUM(E{table_start}:E{total_row - 1})"

    # Second onglet avec une formule inter-onglets, pour vérifier la préservation
    ws2 = wb.create_sheet("Synthese")
    ws2["A1"] = "Total montant"
    ws2["B1"] = f"='{sheet}'!B5"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def make_csv(incoming_dir: Path, day: date, n_rows: int = 5) -> Path:
    df = pd.DataFrame({
        "COLONNE_A_DEFINIR_1": [day.isoformat()] * n_rows,
        "COLONNE_A_DEFINIR_2": ["SITE-TEST"] * n_rows,
        "COLONNE_A_DEFINIR_3": [f"REF-{i:03d}" for i in range(1, n_rows + 1)],
        "COLONNE_A_DEFINIR_4": [f"Article {i}" for i in range(1, n_rows + 1)],
        "COLONNE_A_DEFINIR_5": [10 * i for i in range(1, n_rows + 1)],
        "COLONNE_A_DEFINIR_6": [round(1.5 * i, 2) for i in range(1, n_rows + 1)],
        "COLONNE_IGNOREE": ["x"] * n_rows,
    })
    incoming_dir.mkdir(parents=True, exist_ok=True)
    path = incoming_dir / f"rapport_{day.isoformat()}.csv"
    kw = config.CSV_READ_KWARGS
    df.to_csv(path, index=False, sep=kw.get("sep", ","),
              encoding=kw.get("encoding", "utf-8"), decimal=kw.get("decimal", "."))
    return path


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--template", action="store_true")
    p.add_argument("--csv", action="store_true")
    p.add_argument("--force", action="store_true", help="Écraser un template existant.")
    p.add_argument("--date", type=date.fromisoformat, default=date.today())
    a = p.parse_args(argv)
    do_all = not (a.template or a.csv)

    if a.template or do_all:
        if config.TEMPLATE_PATH.exists() and not a.force:
            print(f"Template déjà présent, non écrasé : {config.TEMPLATE_PATH} (--force pour forcer)")
        else:
            print(f"Template factice créé : {make_template(config.TEMPLATE_PATH)}")
    if a.csv or do_all:
        print(f"CSV factice créé : {make_csv(config.INCOMING_DIR, a.date)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
