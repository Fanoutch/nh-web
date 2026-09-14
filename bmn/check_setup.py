# %% [markdown]
# # Autodiagnostic du projet
#
# À lancer sur le poste de travail, avant et après avoir rempli `config.py`.
# Ne modifie rien : vérifie l'environnement, le template, le mapping et,
# s'il y en a un, le CSV présent dans `data/incoming`.
#
# Usage : `python check_setup.py [--csv chemin.csv]`

# %%
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

OK, KO, WARN = "[OK]  ", "[KO]  ", "[WARN]"
_problems = 0


def report(status: str, msg: str) -> None:
    global _problems
    if status == KO:
        _problems += 1
    print(f"{status} {msg}")


def check_environment() -> None:
    print("\n== Environnement ==")
    v = sys.version_info
    report(OK if v >= (3, 10) else KO, f"Python {v.major}.{v.minor}.{v.micro} (3.10+ requis)")
    for mod in ("pandas", "openpyxl"):
        try:
            m = __import__(mod)
            report(OK, f"{mod} {m.__version__}")
        except ImportError:
            report(KO, f"{mod} absent : pip install -r requirements.txt")


def check_folders(config) -> None:
    print("\n== Dossiers ==")
    for name in ("INCOMING_DIR", "OUTPUT_DIR", "ARCHIVE_DIR", "LOG_DIR"):
        p: Path = getattr(config, name)
        if p.exists():
            report(OK, f"{name} = {p}")
        else:
            report(WARN, f"{name} = {p} (absent, sera créé au premier run)")


def check_template(config):
    print("\n== Template ==")
    if not config.TEMPLATE_PATH.exists():
        report(KO, f"Template introuvable : {config.TEMPLATE_PATH}")
        return None
    from openpyxl import load_workbook
    wb = load_workbook(config.TEMPLATE_PATH, data_only=False)
    report(OK, f"Template chargé : {config.TEMPLATE_PATH.name}, onglets {wb.sheetnames}")
    if config.TARGET_SHEET is None:
        report(WARN, f"TARGET_SHEET = None : premier onglet utilisé ('{wb.sheetnames[0]}')")
        return wb.worksheets[0]
    if config.TARGET_SHEET in wb.sheetnames:
        report(OK, f"Onglet cible '{config.TARGET_SHEET}' présent")
        return wb[config.TARGET_SHEET]
    report(KO, f"Onglet cible '{config.TARGET_SHEET}' absent. Onglets : {wb.sheetnames}")
    return None


def _placeholders(names) -> list[str]:
    return [n for n in names if re.match(r"COLONNE_A_DEFINIR_\d+|A_DEFINIR", n)]


def check_mapping(config, ws) -> None:
    print("\n== Mapping (config.py) ==")
    from daily_report import is_formula_cell
    from openpyxl.cell.cell import MergedCell

    cols = config.required_csv_columns()
    ph = _placeholders(cols)
    if ph:
        report(WARN, f"{len(ph)} placeholder(s) encore présent(s) dans le mapping : {ph}")
    else:
        report(OK, f"Mapping renseigné : {len(cols)} colonnes CSV attendues")

    if ws is None:
        report(WARN, "Template indisponible : vérification des cellules cibles ignorée")
        return

    targets: list[str] = []
    for col, spec in config.CELL_MAPPING.items():
        targets.append(spec if isinstance(spec, str) else spec["cell"])
    tm = config.TABLE_MAPPING
    start, max_rows = tm["start_row"], tm.get("max_rows")
    n_check = max_rows if max_rows else 1
    for letter in tm["columns"].values():
        targets += [f"{letter}{r}" for r in range(start, start + n_check)]

    bad = []
    for coord in targets:
        cell = ws[coord]
        if isinstance(cell, MergedCell):
            bad.append(f"{coord} (fusionnée)")
        elif is_formula_cell(cell):
            bad.append(f"{coord} (formule {cell.value!r})")
    if bad:
        report(KO, f"{len(bad)} cellule(s) cible(s) contiennent une formule ou sont fusionnées : "
                   f"{bad[:8]}{' ...' if len(bad) > 8 else ''}")
    else:
        report(OK, f"{len(targets)} cellules cibles vérifiées, aucune formule ne sera écrasée")

    if max_rows is None:
        report(WARN, "TABLE_MAPPING['max_rows'] = None : aucune protection contre le débordement du tableau")
    else:
        # Ligne juste sous le tableau : signaler si elle contient des formules (pied de tableau)
        below = start + max_rows
        f_below = [f"{l}{below}" for l in tm["columns"].values() if is_formula_cell(ws[f"{l}{below}"])]
        if f_below:
            report(OK, f"Pied de tableau détecté ligne {below} ({f_below[:3]}), protégé par max_rows={max_rows}")


def check_csv(config, csv_path: Path | None) -> None:
    print("\n== CSV ==")
    from daily_report import find_latest_csv, extract_csv_date
    import pandas as pd

    path = csv_path or (find_latest_csv() if config.INCOMING_DIR.exists() else None)
    if path is None:
        report(WARN, f"Aucun CSV dans {config.INCOMING_DIR} (pattern {config.CSV_GLOB}) : lancer "
                     "inspect_csv.py sur le vrai fichier dès réception")
        return
    report(OK, f"CSV candidat : {path.name}")
    d = extract_csv_date(path)
    report(OK if d else WARN,
           f"Date extraite du nom : {d}" if d else
           "Pas de date dans le nom (CSV_DATE_REGEX) : tri par date de modification")
    try:
        df = pd.read_csv(path, **config.CSV_READ_KWARGS)
    except Exception as e:  # noqa: BLE001
        report(KO, f"Lecture impossible avec CSV_READ_KWARGS={config.CSV_READ_KWARGS} : {e}")
        return
    df.columns = [str(c).strip() for c in df.columns]
    if len(df.columns) == 1:
        report(KO, f"Une seule colonne lue ('{df.columns[0][:40]}...') : mauvais séparateur dans CSV_READ_KWARGS ?")
    else:
        report(OK, f"{len(df)} lignes, {len(df.columns)} colonnes : {list(df.columns)}")
    missing = sorted(config.required_csv_columns() - set(df.columns))
    if missing:
        report(KO, f"Colonnes du mapping absentes du CSV : {missing}")
    else:
        report(OK, "Toutes les colonnes du mapping sont présentes dans le CSV")
    mr = config.TABLE_MAPPING.get("max_rows")
    if mr is not None and len(df) > mr:
        report(KO, f"{len(df)} lignes > max_rows={mr}")


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="Diagnostic du projet, sans rien modifier.")
    p.add_argument("--csv", type=Path, default=None)
    a = p.parse_args(argv)

    check_environment()
    if _problems:
        print("\nInstaller les dépendances avant de continuer.")
        return 1
    import config
    check_folders(config)
    ws = check_template(config)
    check_mapping(config, ws)
    check_csv(config, a.csv)

    print("\n== Résultat ==")
    if _problems:
        print(f"{_problems} problème(s) bloquant(s) : corriger config.py / template / CSV puis relancer.")
        return 1
    print("Aucun problème bloquant. Prochaine étape : python daily_report.py --dry-run")
    return 0


if __name__ == "__main__":
    sys.exit(main())
