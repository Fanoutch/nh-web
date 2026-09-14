"""Tests du pipeline avec template et CSV factices, dans un dossier temporaire."""
import json
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

import config
import daily_report as dr
import make_fixtures as mf


@pytest.fixture
def sandbox(tmp_path, monkeypatch):
    """Redirige tous les chemins de config vers tmp_path et y pose les fixtures."""
    monkeypatch.setattr(config, "INCOMING_DIR", tmp_path / "incoming")
    monkeypatch.setattr(config, "OUTPUT_DIR", tmp_path / "output")
    monkeypatch.setattr(config, "ARCHIVE_DIR", tmp_path / "archive")
    monkeypatch.setattr(config, "TEMPLATE_PATH", tmp_path / "template.xlsx")
    monkeypatch.setattr(config, "LOG_FILE", tmp_path / "logs" / "test.log")
    mf.make_template(config.TEMPLATE_PATH)
    csv = mf.make_csv(config.INCOMING_DIR, date(2026, 9, 3))
    return tmp_path, csv


def test_find_latest_csv_prefers_filename_date(sandbox):
    tmp, _ = sandbox
    older = mf.make_csv(config.INCOMING_DIR, date(2026, 9, 1))
    older.touch()  # mtime plus récent, mais date de nom plus ancienne
    latest = dr.find_latest_csv(config.INCOMING_DIR)
    assert latest.name == "rapport_2026-09-03.csv"


def test_find_latest_csv_returns_none_when_empty(tmp_path):
    (tmp_path / "empty").mkdir()
    assert dr.find_latest_csv(tmp_path / "empty") is None


def test_run_returns_no_csv_code_when_incoming_empty(sandbox):
    _, csv = sandbox
    csv.unlink()
    assert dr.run() == config.EXIT_CODE_NO_CSV


def test_end_to_end_preserves_formulas_and_archives(sandbox):
    tmp, csv = sandbox
    ts = datetime(2026, 9, 3, 8, 0, 0)
    res = dr.process(csv, run_ts=ts)
    out = Path(res.output_path)

    assert res.status == "ok" and res.rows_written == 5 and res.cells_written == 3
    assert out.exists()
    assert out.name == "dispo_2026-09-03_20260903_080000.xlsx"
    # CSV archivé, plus dans incoming
    assert not csv.exists()
    archived = list((config.ARCHIVE_DIR / "2026-09-03").glob("*.csv"))
    assert len(archived) == 1

    wb = load_workbook(out, data_only=False)
    ws = wb["Saisie"]
    # Inputs écrits
    assert ws["B2"].value == "2026-09-03"
    assert ws["B3"].value == "SITE-TEST"
    assert ws["B4"].value == 150            # somme des quantités 10..50
    assert ws["A7"].value == "REF-001"
    assert ws["C11"].value == 50
    assert ws["D11"].value == 7.5
    assert ws["A12"].value is None          # ligne suivante non touchée
    # Formules intactes
    assert ws["B5"].value == "=SUM(E7:E56)"
    assert ws["E7"].value == '=IF(C7="","",C7*D7)'
    assert ws["C57"].value == "=SUM(C7:C56)"
    assert wb["Synthese"]["B1"].value == "='Saisie'!B5"
    assert wb.calculation.fullCalcOnLoad is True


def test_dry_run_keeps_csv(sandbox):
    _, csv = sandbox
    dr.process(csv, dry_run=True)
    assert csv.exists()


def test_refuses_to_overwrite_formula(sandbox, monkeypatch):
    _, csv = sandbox
    monkeypatch.setattr(config, "CELL_MAPPING", {"COLONNE_A_DEFINIR_1": "B5"})
    with pytest.raises(dr.FormulaOverwriteError):
        dr.process(csv)
    assert csv.exists()  # rien archivé en cas d'échec


def test_missing_column_fails_before_writing(sandbox, monkeypatch):
    _, csv = sandbox
    monkeypatch.setattr(config, "CELL_MAPPING", {"COLONNE_INEXISTANTE": "B2"})
    with pytest.raises(dr.MissingColumnsError):
        dr.process(csv)
    assert not list(config.OUTPUT_DIR.glob("*.xlsx")) if config.OUTPUT_DIR.exists() else True


def test_too_many_rows_fails(sandbox, monkeypatch):
    _, csv = sandbox
    monkeypatch.setitem(config.TABLE_MAPPING, "max_rows", 3)
    with pytest.raises(dr.PipelineError, match="lignes"):
        dr.process(csv)


def test_run_cli_ok(sandbox):
    assert dr.main([]) == 0
    assert dr.main([]) == config.EXIT_CODE_NO_CSV   # second run : incoming vide


def test_output_dir_override_and_no_archive(sandbox):
    tmp, csv = sandbox
    custom = tmp / "web_out"
    res = dr.process(csv, dry_run=True, output_dir=custom)
    assert Path(res.output_path).parent == custom
    assert csv.exists() and res.archived_to is None


def test_json_output_ok(sandbox, capsys):
    tmp, csv = sandbox
    code = dr.main(["--csv", str(csv), "--no-archive", "--output-dir", str(tmp / "o"), "--json-output"])
    out = capsys.readouterr().out.strip().splitlines()
    data = json.loads(out[-1])
    assert code == 0
    assert data["status"] == "ok" and data["rows_written"] == 5
    assert Path(data["output_path"]).exists()
    assert csv.exists()


def test_json_output_error_is_readable(sandbox, monkeypatch, capsys):
    _, csv = sandbox
    monkeypatch.setattr(config, "CELL_MAPPING", {"COLONNE_INEXISTANTE": "B2"})
    code = dr.main(["--csv", str(csv), "--json-output"])
    data = json.loads(capsys.readouterr().out.strip().splitlines()[-1])
    assert code == config.EXIT_CODE_ERROR
    assert data["status"] == "error" and "COLONNE_INEXISTANTE" in data["message"]


def test_json_output_no_csv(sandbox, capsys):
    _, csv = sandbox
    csv.unlink()
    code = dr.main(["--json-output"])
    data = json.loads(capsys.readouterr().out.strip().splitlines()[-1])
    assert code == config.EXIT_CODE_NO_CSV and data["status"] == "no_csv"
