"""Tests du remplissage par blocs machine (template « dispo »).

Le template de test reproduit la structure du vrai : un bloc par machine,
intitulés en colonne AQ / valeurs en AU, listes HIL et CIL en I / L / R,
et une colonne calculée qui ne doit jamais être écrasée.
"""
import json
import sys
from datetime import date, datetime

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

import config
import daily_report as dr


PAS = 24


def construire_template(chemin, machines=("NH01", "NH02"), intitules=None):
    """Mini template : un bloc de 24 lignes par machine, base à la ligne 9."""
    wb = Workbook()
    ws = wb.active
    ws.title = "DISPO"
    intitules = intitules or {
        "NH01": [(0, "FIA"), (1, "FH"), (2, "ATT"), (3, "TREUI"),
                 (4, "APU OPH"), (5, "APU OPC"), (10, "FIAA")],
        "NH02": [(0, "FIA"), (1, "FH"), (2, "ATT"), (3, "TREUI"),
                 (4, "APU OPH"), (5, "APU OPC"), (14, "Cable RHA")],
    }
    for i, machine in enumerate(machines):
        base = 9 + i * PAS
        ws[f"A{base}"] = machine
        for decalage, libelle in intitules.get(machine, []):
            ws[f"AQ{base + decalage}"] = libelle
        ws[f"F{base + 6}"] = "HIL"
        ws[f"F{base + 12}"] = "CIL"
        ws[f"AJ{base + 1}"] = f"=AU{base + 1}*2"   # colonne calculée, à protéger
    wb.save(chemin)
    return chemin


@pytest.fixture
def blocs_actifs(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "TEMPLATE_PATH", construire_template(tmp_path / "dispo.xlsx"))
    monkeypatch.setattr(config, "OUTPUT_DIR", tmp_path / "out")
    monkeypatch.setattr(config, "LOG_FILE", tmp_path / "logs" / "t.log")
    monkeypatch.setattr(config, "BLOCS", dict(config.BLOCS, actif=True))
    return tmp_path


def donnees(lignes):
    return pd.DataFrame(lignes)


def feuille(chemin):
    return load_workbook(chemin, data_only=False)["DISPO"]


# ---------------------------------------------------------------------------
# Repérage des blocs
# ---------------------------------------------------------------------------
def test_reperer_les_blocs_par_le_nom_de_machine(blocs_actifs):
    wb = load_workbook(config.TEMPLATE_PATH)
    assert dr.reperer_blocs(wb["DISPO"]) == {"NH01": 9, "NH02": 33}


# ---------------------------------------------------------------------------
# Valeurs repérées par intitulé
# ---------------------------------------------------------------------------
def test_ecrit_les_compteurs_en_face_de_leur_intitule(blocs_actifs):
    df = donnees([{"machine": "NH01", "fh": 2581.73, "att": 4357, "treuil": 4591,
                   "apu_oph": 169.41, "apu_opc": 1043}])
    wb = load_workbook(config.TEMPLATE_PATH)
    dr.fill_blocs(wb["DISPO"], df)

    ws = wb["DISPO"]
    assert ws["AU10"].value == 2581.73     # FH, en face de AQ10
    assert ws["AU11"].value == 4357        # ATT
    assert ws["AU14"].value == 1043        # APU OPC
    assert ws["AU33"].value is None        # NH02 n'a pas de données


def test_intitule_absent_du_bloc_est_ignore_sans_planter(blocs_actifs, caplog):
    df = donnees([{"machine": "NH02", "fh": 100, "fiaa": 42}])
    wb = load_workbook(config.TEMPLATE_PATH)
    config.BLOCS["valeurs"]["champs"]["fiaa"] = "FIAA"
    try:
        dr.fill_blocs(wb["DISPO"], df)
    finally:
        config.BLOCS["valeurs"]["champs"].pop("fiaa")

    assert wb["DISPO"]["AU34"].value == 100        # FH écrit
    assert "FIAA" in caplog.text                   # l'absence est signalée


def test_nentoure_jamais_une_formule(blocs_actifs):
    df = donnees([{"machine": "NH01", "fh": 1}])
    wb = load_workbook(config.TEMPLATE_PATH)
    dr.fill_blocs(wb["DISPO"], df)
    assert wb["DISPO"]["AJ10"].value == "=AU10*2"


# ---------------------------------------------------------------------------
# Listes HIL / CIL
# ---------------------------------------------------------------------------
def test_ecrit_les_lignes_hil_et_cil(blocs_actifs):
    df = donnees([
        {"machine": "NH01", "fh": 10},
        {"machine": "NH01", "zone": "hil", "date": date(2026, 3, 27),
         "equipement": "PDU (P)", "ref": "M31900000814"},
        {"machine": "NH01", "zone": "cil", "date": date(2026, 4, 1),
         "equipement": "Empreinte vis CCU", "ref": "MH3102907356"},
    ])
    wb = load_workbook(config.TEMPLATE_PATH)
    dr.fill_blocs(wb["DISPO"], df)

    ws = wb["DISPO"]
    assert ws["I15"].value == date(2026, 3, 27)     # HIL : premier emplacement
    assert ws["L15"].value == "PDU (P)"
    assert ws["R15"].value == "M31900000814"
    assert ws["I21"].value == date(2026, 4, 1)      # CIL : premier emplacement
    assert ws["L21"].value == "Empreinte vis CCU"


def test_vide_les_anciennes_lignes_avant_decrire(blocs_actifs):
    wb = load_workbook(config.TEMPLATE_PATH)
    ws = wb["DISPO"]
    ws["L15"], ws["L16"] = "ancien 1", "ancien 2"

    df = donnees([{"machine": "NH01", "zone": "hil", "date": date(2026, 1, 1),
                   "equipement": "nouveau", "ref": "X1"}])
    dr.fill_blocs(ws, df)

    assert ws["L15"].value == "nouveau"
    assert ws["L16"].value is None


def test_trop_de_lignes_pour_la_place_disponible(blocs_actifs, caplog):
    lignes = [{"machine": "NH01", "zone": "hil", "date": date(2026, 1, i + 1),
               "equipement": f"eq{i}", "ref": f"R{i}"} for i in range(8)]
    wb = load_workbook(config.TEMPLATE_PATH)
    dr.fill_blocs(wb["DISPO"], donnees(lignes))

    ws = wb["DISPO"]
    assert ws["L20"].value == "eq5"          # 6 emplacements remplis
    assert "2" in caplog.text and "HIL" in caplog.text.upper()   # 2 laissées de côté


def test_machine_absente_du_template_nempeche_pas_les_autres(blocs_actifs, caplog):
    df = donnees([{"machine": "NH99", "fh": 1}, {"machine": "NH01", "fh": 2}])
    wb = load_workbook(config.TEMPLATE_PATH)
    dr.fill_blocs(wb["DISPO"], df)

    assert wb["DISPO"]["AU10"].value == 2
    assert "NH99" in caplog.text


def test_bloc_sans_donnees_reste_intact(blocs_actifs):
    wb = load_workbook(config.TEMPLATE_PATH)
    ws = wb["DISPO"]
    ws["L39"] = "déjà là"
    dr.fill_blocs(ws, donnees([{"machine": "NH01", "fh": 1}]))
    assert ws["L39"].value == "déjà là"


# ---------------------------------------------------------------------------
# Bout en bout
# ---------------------------------------------------------------------------
def test_process_ecrit_le_template_par_blocs(blocs_actifs, tmp_path):
    csv = tmp_path / "rapport_2026-09-03.csv"
    csv.write_text(
        "machine;zone;date;equipement;ref;fh;att\n"
        "NH01;;;;;2581,73;4357\n"
        "NH01;hil;2026-03-27;PDU (P);M31900000814;;\n"
        "NH02;;;;;2339,16;4716\n",
        encoding="utf-8-sig")

    resultat = dr.process(csv, dry_run=True, output_dir=tmp_path / "out")

    assert resultat.status == "ok"
    ws = feuille(resultat.output_path)
    assert ws["AU10"].value == 2581.73
    assert ws["L15"].value == "PDU (P)"
    assert ws["I15"].value == datetime(2026, 3, 27)   # vraie date, comme en JSON
    assert ws["AU34"].value == 2339.16


# ---------------------------------------------------------------------------
# Source JSON
# ---------------------------------------------------------------------------
def test_json_imbrique_devient_des_lignes(blocs_actifs, tmp_path):
    fichier = tmp_path / "dispo_2026-09-03.json"
    fichier.write_text(json.dumps([{
        "machine": "NH01", "fh": 2581.73, "att": 4357,
        "hil": [{"date": "2026-03-27", "equipement": "PDU (P)", "ref": "M319"}],
        "cil": [{"date": "2026-04-01", "equipement": "Vis CCU", "ref": "MH310"}],
    }]), encoding="utf-8")

    df = dr.load_source(fichier)

    assert list(df["machine"]) == ["NH01"] * 3
    assert set(df["zone"].dropna()) == {"hil", "cil"}
    assert df.loc[df.zone == "hil", "equipement"].iloc[0] == "PDU (P)"


def test_json_liste_sous_une_cle(blocs_actifs, tmp_path):
    fichier = tmp_path / "dispo_2026-09-03.json"
    fichier.write_text(json.dumps({"export": "x", "machines": [{"machine": "NH02", "fh": 12}]}),
                       encoding="utf-8")
    assert dr.load_source(fichier)["fh"].iloc[0] == 12


def test_process_accepte_un_json(blocs_actifs, tmp_path):
    fichier = tmp_path / "dispo_2026-09-03.json"
    fichier.write_text(json.dumps([{
        "machine": "NH01", "fh": 2581.73,
        "hil": [{"date": "2026-03-27", "equipement": "PDU (P)", "ref": "M319"}],
    }]), encoding="utf-8")

    resultat = dr.process(fichier, dry_run=True, output_dir=tmp_path / "out")

    ws = feuille(resultat.output_path)
    assert ws["AU10"].value == 2581.73
    assert ws["L15"].value == "PDU (P)"
    assert ws["I15"].value == datetime(2026, 3, 27)     # date, pas texte


# ---------------------------------------------------------------------------
# Libellé HIL produit par le LLM
# ---------------------------------------------------------------------------
def test_colonne_avec_repli_prend_le_llm_si_la_source_est_vide(blocs_actifs):
    df = donnees([
        {"machine": "NH01", "zone": "hil", "date": date(2026, 1, 1),
         "equipement": None, "llm.equipement": "PDU (P)", "ref": "X1"},
        {"machine": "NH01", "zone": "hil", "date": date(2026, 1, 2),
         "equipement": "saisi à la main", "llm.equipement": "ignoré", "ref": "X2"},
    ])
    wb = load_workbook(config.TEMPLATE_PATH)
    dr.fill_blocs(wb["DISPO"], df)

    ws = wb["DISPO"]
    assert ws["L15"].value == "PDU (P)"           # repli sur la réponse du modèle
    assert ws["L16"].value == "saisi à la main"   # la source prime


def test_le_llm_nest_interroge_que_sur_les_lignes_hil(blocs_actifs, monkeypatch):
    appels = []

    class FauxClient:
        LlmError = RuntimeError

        @staticmethod
        def extraire_equipement(enregistrement, conf):
            appels.append(enregistrement)
            return {"equipement": "PDU (P)", "indice": "haut", "justification": "x"}

    monkeypatch.setitem(sys.modules, "llm_client", FauxClient)
    monkeypatch.setattr(config, "LLM", dict(config.LLM, actif=True, hors_ligne=False,
                                            filtre={"champ": "zone", "valeurs": ["hil"]}))
    df = donnees([
        {"machine": "NH01", "zone": None, "fh": 10},
        {"machine": "NH01", "zone": "hil", "travail_demande": "vibration PDU",
         "travail_effectue": "dépose PDU"},
        {"machine": "NH01", "zone": "cil", "travail_demande": "autre"},
    ])
    resultat = dr.enrichir_avec_llm(df)

    assert len(appels) == 1                                   # une seule ligne envoyée
    assert "vibration PDU" in str(appels[0])
    assert list(resultat["llm.equipement"]) == ["", "PDU (P)", ""]


def test_le_prompt_reprend_les_exemples_de_style(blocs_actifs):
    import llm_client
    conf = dict(config.LLM, exemples=["PDU (P)", "RHEAS"])
    messages = llm_client.construire_messages(
        {"travail_demande": "x", "travail_effectue": "y"}, conf)
    assert "PDU (P)" in messages[0]["content"] and "RHEAS" in messages[0]["content"]
    assert "{exemples}" not in messages[0]["content"]
    assert "SAISIES HUMAINES" in messages[0]["content"]      # le modèle est prévenu
    assert "NOM DE LA PIÈCE" in messages[0]["content"]       # et sait quoi rendre
