"""Suppression des blocs des NH absentes de l'extraction.

Reproduit la suppression de lignes d'Excel : cellules remontées, formules,
fusions, mises en forme conditionnelles, hauteurs de ligne et zone d'impression
décalées ; une référence vers une ligne supprimée devient #REF!, comme dans Excel.
"""
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill

import config
import daily_report as dr
import suppression_blocs as sb

PAS = 24


def template_trois_blocs(chemin):
    """En-tête (lignes 1-8) + NH01 (9) + emplacement sans nom (33) + NH03 (57) + pied (81)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "DISPO"
    ws["C2"] = "Etat de disponibilité du"
    ws["BC4"] = None
    for base, nom in ((9, "NH01"), (33, None), (57, "NH03")):
        if nom:
            ws[f"A{base}"] = nom
        ws[f"AQ{base + 1}"] = "FH"
        ws[f"R{base}"] = f"=BG{base + 8}"                       # même bloc
        ws[f"AJ{base + 9}"] = f'=IF(ISNUMBER(W{base + 9}),(W{base + 9}-BC4),"")'  # en-tête
        ws.merge_cells(f"L{base + 6}:Q{base + 6}")
        ws.row_dimensions[base + 1].height = 20 + base          # hauteur propre au bloc
    ws["L15"] = "HIL NH01"
    ws["L63"] = "HIL NH03"
    ws["AU58"] = 2339.16
    ws["AK66"] = '=IF(ISNUMBER(W66),(W66-BC29),"")'             # pointe dans le bloc NH01
    ws["AQ58"].comment = Comment("note NH03", "test")
    ws["A81"] = "<<<<<<<<<<<"
    ws.conditional_formatting.add(
        "I15:K17 I63:K65",
        CellIsRule(operator="lessThan", formula=["$BC$4+15"], fill=PatternFill("solid", fgColor="FF0000")))
    ws.print_area = "A1:AM80"
    wb.save(chemin)
    return chemin


@pytest.fixture
def feuille(tmp_path):
    return load_workbook(template_trois_blocs(tmp_path / "dispo.xlsx"))["DISPO"]


# ---------------------------------------------------------------------------
# Choix des blocs à supprimer
# ---------------------------------------------------------------------------
def test_blocs_absents_et_emplacement_sans_nom(feuille):
    lignes = sb.lignes_des_blocs_absents(feuille, presentes={"NH03"}, hauteur=PAS)
    assert lignes == set(range(9, 57))          # NH01 + emplacement sans nom


def test_aucune_suppression_si_toutes_presentes(feuille):
    assert sb.lignes_des_blocs_absents(feuille, presentes={"NH01", "NH03"}, hauteur=PAS) \
        == set(range(33, 57))                   # seul l'emplacement sans nom part


# ---------------------------------------------------------------------------
# Réécriture de formule (règles d'Excel)
# ---------------------------------------------------------------------------
def carte(supprimees, max_row=200):
    return sb.table_de_decalage(set(supprimees), max_row)


def test_formule_decalee_vers_le_haut():
    m = carte(range(9, 33))
    assert sb.decaler_formule("=BG65", m, "DISPO") == "=BG41"


def test_reference_en_tete_inchangee():
    m = carte(range(9, 33))
    assert sb.decaler_formule('=IF(ISNUMBER(W66),(W66-BC4),"")', m, "DISPO") \
        == '=IF(ISNUMBER(W42),(W42-BC4),"")'


def test_reference_vers_ligne_supprimee_devient_ref():
    m = carte(range(9, 33))
    assert sb.decaler_formule("=W66-BC29", m, "DISPO") == "=W42-#REF!"


def test_references_absolues_et_texte_preserves():
    m = carte(range(9, 33))
    assert sb.decaler_formule('=$BC$40&"A40"', m, "DISPO") == '=$BC$16&"A40"'


def test_autre_feuille_non_decalee():
    m = carte(range(9, 33))
    assert sb.decaler_formule("=Tuto!A40+A40", m, "DISPO") == "=Tuto!A40+A16"


def test_plage_retrecie_ou_supprimee():
    m = carte(range(9, 33))
    assert sb.decaler_formule("=SUM(A5:A40)", m, "DISPO") == "=SUM(A5:A16)"
    assert sb.decaler_formule("=SUM(A10:A20)", m, "DISPO") == "=SUM(#REF!)"


# ---------------------------------------------------------------------------
# Suppression complète sur une feuille
# ---------------------------------------------------------------------------
def test_suppression_remonte_les_cellules_et_les_formules(feuille):
    sb.supprimer_lignes(feuille, set(range(9, 57)))

    assert feuille["A9"].value == "NH03"
    assert feuille["L15"].value == "HIL NH03"
    assert feuille["AU10"].value == 2339.16
    assert feuille["R9"].value == "=BG17"                    # était =BG65 en ligne 57
    assert feuille["AJ18"].value == '=IF(ISNUMBER(W18),(W18-BC4),"")'   # en-tête conservé
    assert feuille["A33"].value == "<<<<<<<<<<<"             # pied remonté de 48 lignes
    assert feuille.max_row == 33


def test_aucune_cellule_des_lignes_supprimees_ne_subsiste(feuille):
    sb.supprimer_lignes(feuille, set(range(9, 57)))
    valeurs = [c.value for row in feuille.iter_rows() for c in row if c.value is not None]
    assert "HIL NH01" not in valeurs and "NH01" not in valeurs


def test_reference_croisee_vers_bloc_supprime_devient_ref(feuille):
    """AK66 pointe vers BC29, dans le bloc NH01 supprimé : #REF!, comme Excel."""
    sb.supprimer_lignes(feuille, set(range(9, 57)))
    assert feuille["AK18"].value == '=IF(ISNUMBER(W18),(W18-#REF!),"")'


def test_fusions_decalees_et_supprimees(feuille):
    sb.supprimer_lignes(feuille, set(range(9, 57)))
    fusions = {str(r) for r in feuille.merged_cells.ranges}
    assert fusions == {"L15:Q15"}                            # celle de NH03 (L63), remontée


def test_mise_en_forme_conditionnelle_decalee(feuille):
    sb.supprimer_lignes(feuille, set(range(9, 57)))
    plages = [str(cf.sqref) for cf in feuille.conditional_formatting]
    assert plages == ["I15:K17"]                             # I63:K65 remontée, I15:K17 (NH01) supprimée
    regle = list(feuille.conditional_formatting)[0].rules[0]
    assert regle.formula == ["$BC$4+15"]


def test_hauteurs_de_ligne_decalees(feuille):
    sb.supprimer_lignes(feuille, set(range(9, 57)))
    assert feuille.row_dimensions[10].height == 20 + 57      # hauteur de la ligne 58


def test_commentaire_suit_sa_cellule(feuille):
    sb.supprimer_lignes(feuille, set(range(9, 57)))
    assert feuille["AQ10"].comment is not None
    assert feuille["AQ10"].comment.text == "note NH03"


def test_zone_dimpression_ajustee(feuille):
    sb.supprimer_lignes(feuille, set(range(9, 57)))
    assert feuille.print_area.endswith("$AM$32")


# ---------------------------------------------------------------------------
# Intégration dans le pipeline
# ---------------------------------------------------------------------------
@pytest.fixture
def pipeline(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "TEMPLATE_PATH", template_trois_blocs(tmp_path / "dispo.xlsx"))
    monkeypatch.setattr(config, "LOG_FILE", tmp_path / "logs" / "t.log")
    monkeypatch.setattr(config, "BLOCS", dict(config.BLOCS, actif=True, supprimer_absentes=True))
    return tmp_path


def test_process_ne_garde_que_les_nh_de_lextraction(pipeline):
    source = pipeline / "dispo_2026-09-16.csv"
    source.write_text("machine;zone;date;equipement;ref;fh\nNH03;;;;;1500\n", encoding="utf-8-sig")

    resultat = dr.process(source, dry_run=True, output_dir=pipeline / "out")

    ws = load_workbook(resultat.output_path)["DISPO"]
    noms = [c.value for c in ws["A"] if isinstance(c.value, str) and c.value.startswith("NH")]
    assert noms == ["NH03"]
    assert ws["A9"].value == "NH03" and ws["AU10"].value == 1500


def test_option_desactivee_garde_tous_les_blocs(pipeline, monkeypatch):
    monkeypatch.setattr(config, "BLOCS", dict(config.BLOCS, supprimer_absentes=False))
    source = pipeline / "dispo_2026-09-16.csv"
    source.write_text("machine;zone;date;equipement;ref;fh\nNH03;;;;;1500\n", encoding="utf-8-sig")

    resultat = dr.process(source, dry_run=True, output_dir=pipeline / "out")

    ws = load_workbook(resultat.output_path)["DISPO"]
    assert ws["A9"].value == "NH01" and ws["A57"].value == "NH03"


# ---------------------------------------------------------------------------
# Vrai template (poste de dev uniquement : ignoré par git)
# ---------------------------------------------------------------------------
VRAI = Path(config.BASE_DIR) / "template" / "dispo.xlsx"


@pytest.mark.skipif(not VRAI.exists(), reason="template réel absent (bmn/template/dispo.xlsx)")
def test_vrai_template_formules_coherentes_apres_suppression():
    wb = load_workbook(VRAI)
    ws = wb["DISPO"]
    avant = load_workbook(VRAI)["DISPO"]
    blocs = dr.reperer_blocs(avant)
    garder = {"NH03", "NH10", "NH27"}
    lignes = sb.lignes_des_blocs_absents(ws, presentes=garder, hauteur=PAS)
    table = sb.table_de_decalage(lignes, ws.max_row)
    sb.supprimer_lignes(ws, lignes)

    # chaque formule d'un bloc conservé = formule d'origine décalée selon les règles d'Excel
    for nom in garder:
        base = blocs[nom]
        for r in range(base, base + PAS):
            for c in avant[r]:
                if isinstance(c.value, str) and c.value.startswith("="):
                    attendu = sb.decaler_formule(c.value, table, "DISPO")
                    assert ws.cell(table[r], c.column).value == attendu, c.coordinate
    noms = [c.value.strip() for c in ws["A"] if isinstance(c.value, str) and c.value.strip().startswith("NH")]
    assert noms == ["NH03", "NH10", "NH27"]
    assert all(r.max_row <= ws.max_row for r in ws.merged_cells.ranges)
