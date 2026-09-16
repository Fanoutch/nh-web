# %% [markdown]
# # Génération quotidienne Excel à partir d'un CSV
#
# Pipeline : repérer le CSV du jour -> charger le template (formules préservées)
# -> écrire uniquement les cellules d'input -> sauvegarder horodaté -> archiver le CSV.
#
# Exécution : `python daily_report.py` (voir `--help`).
# Les cellules `# %%` permettent de tester chaque étape dans VS Code.

# %%
from __future__ import annotations

import argparse
import json
import logging
import re
import shutil
import sys
from dataclasses import asdict, dataclass
from datetime import date, datetime
from logging.handlers import RotatingFileHandler
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
from openpyxl.worksheet.worksheet import Worksheet

import config

log = logging.getLogger("daily_report")


# ---------------------------------------------------------------------------
# Exceptions
# ---------------------------------------------------------------------------
class PipelineError(Exception):
    """Erreur fonctionnelle : le run doit s'arrêter sans produire de sortie."""


class FormulaOverwriteError(PipelineError):
    """Tentative d'écriture dans une cellule contenant une formule."""


class MissingColumnsError(PipelineError):
    """Le CSV ne contient pas toutes les colonnes attendues par le mapping."""


# ---------------------------------------------------------------------------
# Logging
# ---------------------------------------------------------------------------
def setup_logging(log_file: Path | None = None, level: int = logging.INFO,
                  console_stream=None) -> None:
    """Journalise dans un fichier rotatif + sur la console. Idempotent.

    `console_stream` : sys.stdout par défaut ; passer sys.stderr en mode
    --json-output pour garder stdout propre.
    """
    if getattr(setup_logging, "_configured", False):
        return
    log_file = log_file or config.LOG_FILE
    log_file.parent.mkdir(parents=True, exist_ok=True)
    fmt = logging.Formatter("%(asctime)s | %(levelname)-7s | %(message)s")

    fh = RotatingFileHandler(
        log_file, maxBytes=config.LOG_MAX_BYTES,
        backupCount=config.LOG_BACKUP_COUNT, encoding="utf-8",
    )
    fh.setFormatter(fmt)
    sh = logging.StreamHandler(console_stream or sys.stdout)
    sh.setFormatter(fmt)

    log.setLevel(level)
    log.addHandler(fh)
    log.addHandler(sh)
    setup_logging._configured = True  # type: ignore[attr-defined]


# ---------------------------------------------------------------------------
# 1. Repérage du CSV du jour
# ---------------------------------------------------------------------------
def extract_csv_date(path: Path) -> date | None:
    """Date lue dans le nom du fichier via config.CSV_DATE_REGEX, sinon None."""
    if not config.CSV_DATE_REGEX:
        return None
    m = re.search(config.CSV_DATE_REGEX, path.name)
    if not m:
        return None
    try:
        return datetime.strptime(m.group("date"), config.CSV_DATE_FORMAT).date()
    except ValueError:
        return None


def csv_sort_key(path: Path) -> tuple[date, float]:
    """Tri : date du nom de fichier en priorité, puis date de modification."""
    d = extract_csv_date(path) or date.fromtimestamp(path.stat().st_mtime)
    return (d, path.stat().st_mtime)


def find_latest_csv(incoming_dir: Path | None = None,
                    pattern: str | None = None) -> Path | None:
    """Retourne le CSV le plus récent du dossier entrant, ou None s'il n'y en a pas."""
    incoming_dir = incoming_dir or config.INCOMING_DIR
    pattern = pattern or config.CSV_GLOB
    candidates = [p for p in incoming_dir.glob(pattern) if p.is_file()]
    if not candidates:
        return None
    candidates.sort(key=csv_sort_key)
    if len(candidates) > 1:
        log.warning("%d CSV présents dans %s, le plus récent est retenu : %s",
                    len(candidates), incoming_dir, candidates[-1].name)
    return candidates[-1]


# ---------------------------------------------------------------------------
# 2. Chargement et validation du CSV
# ---------------------------------------------------------------------------
def load_csv(path: Path) -> pd.DataFrame:
    df = pd.read_csv(path, **config.CSV_READ_KWARGS)
    df.columns = [str(c).strip() for c in df.columns]
    log.info("CSV chargé : %s (%d lignes, %d colonnes)", path.name, len(df), len(df.columns))
    return df


def load_source(path: Path) -> pd.DataFrame:
    """Charge la source du jour : CSV, ou JSON (plat, imbriqué, ou JSON Lines)."""
    if path.suffix.lower() in (".json", ".jsonl"):
        return load_json(path)
    return _convertir_dates(load_csv(path))


def load_json(path: Path) -> pd.DataFrame:
    """Aplati un JSON en lignes : une ligne de compteurs par machine, puis une
    ligne par entrée de liste (HIL, CIL), repérée par la colonne « zone »."""
    texte = path.read_text(encoding="utf-8-sig").strip()
    try:
        donnees = json.loads(texte)
    except json.JSONDecodeError as exc:
        try:
            donnees = [json.loads(l) for l in texte.splitlines() if l.strip()]
        except json.JSONDecodeError:
            raise PipelineError(f"JSON illisible ({exc.msg}, ligne {exc.lineno}) : {path.name}")

    enregistrements = _enregistrements_json(donnees)
    conf = config.BLOCS
    champ_machine, champ_zone = conf["champ_machine"], conf["champ_zone"]
    zones = list(conf["listes"])

    lignes: list[dict] = []
    for enr in enregistrements:
        if not isinstance(enr, dict):
            continue
        lignes.append({cle: val for cle, val in enr.items()
                       if cle not in zones and not isinstance(val, (list, dict))})
        for zone in zones:
            for entree in enr.get(zone) or []:
                if not isinstance(entree, dict):
                    continue
                ligne = {champ_machine: enr.get(champ_machine), champ_zone: zone}
                ligne.update({c: v for c, v in entree.items()
                              if not isinstance(v, (list, dict))})
                lignes.append(ligne)

    df = _convertir_dates(pd.DataFrame(lignes))
    log.info("JSON chargé : %s (%d enregistrement(s) -> %d ligne(s))",
             path.name, len(enregistrements), len(df))
    return df


def _enregistrements_json(donnees):
    """Retrouve la liste d'enregistrements : racine, ou plus longue liste d'objets."""
    if isinstance(donnees, list):
        return donnees
    if isinstance(donnees, dict):
        listes = [v for v in donnees.values()
                  if isinstance(v, list) and v and isinstance(v[0], dict)]
        return max(listes, key=len) if listes else [donnees]
    raise PipelineError("JSON inexploitable : ni liste, ni objet.")


def _convertir_dates(df: pd.DataFrame) -> pd.DataFrame:
    """Les champs de date deviennent de vraies dates (sinon Excel affiche du texte)."""
    for champ in config.BLOCS.get("champs_dates", []):
        if champ in df.columns:
            df[champ] = pd.to_datetime(df[champ], errors="coerce", format="mixed")
    return df


def validate_columns(df: pd.DataFrame,
                     required: set[str] | None = None) -> None:
    required = required if required is not None else config.required_csv_columns()
    missing = sorted(required - set(df.columns))
    if missing:
        raise MissingColumnsError(
            f"Colonnes manquantes dans le CSV : {missing}. "
            f"Colonnes disponibles : {list(df.columns)}"
        )


# ---------------------------------------------------------------------------
# 3. Template Excel
# ---------------------------------------------------------------------------
# ---------------------------------------------------------------------------
# 4 bis. Remplissage par blocs machine (template « dispo »)
# ---------------------------------------------------------------------------
def reperer_blocs(ws: Worksheet, conf: dict | None = None) -> dict[str, int]:
    """Nom de machine -> première ligne de son bloc, lu en colonne A du template."""
    conf = conf or config.BLOCS
    motif = re.compile(conf["motif_machine"], re.IGNORECASE)
    colonne = conf["colonne_machine"]
    blocs: dict[str, int] = {}
    for ligne in range(1, ws.max_row + 1):
        valeur = ws[f"{colonne}{ligne}"].value
        if isinstance(valeur, str) and motif.fullmatch(valeur.strip()):
            blocs[normaliser_machine(valeur)] = ligne
    return blocs


def normaliser_machine(nom) -> str:
    """« nh 01 » et « NH01 » désignent la même machine."""
    return str(nom).replace(" ", "").upper()


def trouver_ligne_intitule(ws: Worksheet, base: int, intitule: str,
                           conf: dict) -> int | None:
    """Ligne du bloc dont la colonne d'intitulés porte ce libellé."""
    colonne = conf["valeurs"]["colonne_libelle"]
    attendu = intitule.strip().casefold()
    for ligne in range(base, base + config_pas_bloc(ws, base)):
        valeur = ws[f"{colonne}{ligne}"].value
        if isinstance(valeur, str) and valeur.strip().casefold() == attendu:
            return ligne
    return None


def config_pas_bloc(ws: Worksheet, base: int) -> int:
    """Hauteur d'un bloc : jusqu'au bloc suivant, sinon jusqu'au bas de la feuille."""
    suivantes = [l for l in _lignes_blocs(ws) if l > base]
    return (min(suivantes) - base) if suivantes else (ws.max_row - base + 1)


def _lignes_blocs(ws: Worksheet) -> list[int]:
    if not hasattr(ws, "_bmn_lignes_blocs"):
        ws._bmn_lignes_blocs = sorted(reperer_blocs(ws).values())
    return ws._bmn_lignes_blocs


def _valeur_renseignee(valeur) -> bool:
    return not (valeur is None or (isinstance(valeur, float) and pd.isna(valeur))
                or str(valeur).strip() == "")


def _cible_colonne(champ: str, reglage) -> tuple[str, list[str]]:
    """Une colonne de liste : « champ: "L" », ou « champ: {colonne, champs} ».

    La seconde forme permet un repli : on prend le premier champ renseigné,
    par exemple le libellé de la source, sinon celui produit par le LLM.
    """
    if isinstance(reglage, dict):
        return reglage["colonne"], list(reglage.get("champs") or [champ])
    return reglage, [champ]


def _premiere_valeur(entree: pd.Series, champs: list[str]):
    for champ in champs:
        if champ in entree.index and _valeur_renseignee(entree[champ]):
            return entree[champ]
    return None


def fill_blocs(ws: Worksheet, df: pd.DataFrame, conf: dict | None = None) -> tuple[int, int]:
    """Remplit un bloc par machine. Retourne (cellules écrites, lignes de liste écrites)."""
    conf = conf or config.BLOCS
    blocs = reperer_blocs(ws, conf)
    champ_machine, champ_zone = conf["champ_machine"], conf["champ_zone"]
    if champ_machine not in df.columns:
        raise MissingColumnsError(
            f"Colonne « {champ_machine} » absente de la source : impossible de savoir "
            "à quelle machine rattacher les données.")

    n_cellules = n_lignes = 0
    for machine, lignes in df.groupby(df[champ_machine].map(normaliser_machine), sort=False):
        base = blocs.get(machine)
        if base is None:
            log.warning("Machine %s absente du template (%d ligne(s) ignorée(s))",
                        machine, len(lignes))
            continue
        n_cellules += _ecrire_valeurs(ws, base, machine, lignes, conf, champ_zone)
        n_lignes += _ecrire_listes(ws, base, machine, lignes, conf, champ_zone)

    log.info("Blocs remplis : %d cellule(s), %d ligne(s) de liste", n_cellules, n_lignes)
    return n_cellules, n_lignes


def _ecrire_valeurs(ws: Worksheet, base: int, machine: str, lignes: pd.DataFrame,
                    conf: dict, champ_zone: str) -> int:
    """Compteurs repérés par leur intitulé (colonne AQ -> colonne AU)."""
    colonne_valeur = conf["valeurs"]["colonne_valeur"]
    ecrites = 0
    for champ, intitule in conf["valeurs"]["champs"].items():
        if champ not in lignes.columns:
            continue
        valeurs = [v for v in lignes[champ] if _valeur_renseignee(v)]
        if not valeurs:
            continue
        ligne = trouver_ligne_intitule(ws, base, intitule, conf)
        if ligne is None:
            log.warning("Machine %s : intitulé « %s » absent de son bloc, %s non écrit",
                        machine, intitule, champ)
            continue
        write_value(ws, f"{colonne_valeur}{ligne}", valeurs[0])
        ecrites += 1
    return ecrites


def _ecrire_listes(ws: Worksheet, base: int, machine: str, lignes: pd.DataFrame,
                   conf: dict, champ_zone: str) -> int:
    """Listes HIL / CIL : n emplacements consécutifs, plusieurs colonnes par ligne."""
    ecrites = 0
    zones = lignes[champ_zone] if champ_zone in lignes.columns else None
    for nom_zone, reglage in conf["listes"].items():
        if zones is None:
            entrees = lignes.iloc[0:0]
        else:
            entrees = lignes[zones.map(
                lambda v: _valeur_renseignee(v) and str(v).strip().casefold() == nom_zone)]
        premier = base + reglage["decalage"]
        places = reglage["emplacements"]

        cibles = [_cible_colonne(champ, reglage_colonne)
                  for champ, reglage_colonne in reglage["colonnes"].items()]

        if conf.get("vider_avant_ecriture") and len(entrees):
            for i in range(places):
                for colonne, _ in cibles:
                    write_value(ws, f"{colonne}{premier + i}", None)

        for i, (_, entree) in enumerate(entrees.iterrows()):
            if i >= places:
                break
            for colonne, champs in cibles:
                valeur = _premiere_valeur(entree, champs)
                if valeur is not None:
                    write_value(ws, f"{colonne}{premier + i}", valeur)
            ecrites += 1

        if len(entrees) > places:
            log.warning("Machine %s : %d ligne(s) %s laissée(s) de côté, le template "
                        "n'a que %d emplacement(s)",
                        machine, len(entrees) - places, nom_zone.upper(), places)
    return ecrites


def _lignes_a_interroger(df: pd.DataFrame, conf: dict) -> pd.DataFrame:
    """Sous-ensemble des lignes envoyées au modèle (filtre de config.LLM).

    Sans filtre, toutes les lignes. Avec `{"champ": "zone", "valeurs": ["hil"]}`,
    seules les lignes HIL, qui seules ont besoin d'un libellé d'équipement.
    """
    filtre = conf.get("filtre") or {}
    champ = filtre.get("champ")
    if not champ or champ not in df.columns:
        return df
    attendues = {str(v).strip().casefold() for v in filtre.get("valeurs", [])}
    garder = df[champ].map(
        lambda v: _valeur_renseignee(v) and str(v).strip().casefold() in attendues)
    return df[garder]


def enrichir_avec_llm(df: pd.DataFrame) -> pd.DataFrame:
    """Ajoute les colonnes `llm.*` au tableau (vides si le LLM est inactif).

    Une ligne dont l'appel échoue reste vide : la dispo est générée quand même,
    l'incident est journalisé. Les colonnes existent toujours, pour qu'un
    mapping qui les référence ne casse pas quand le LLM est désactivé.
    """
    conf = config.LLM
    colonnes = [f"{config.LLM_PREFIXE}{champ}" for champ in conf["champs_produits"]]
    for colonne in colonnes:
        if colonne not in df.columns:
            df[colonne] = ""

    if not conf.get("actif"):
        return df

    import llm_client  # import tardif : le pipeline tourne sans lui si inactif

    concernees = _lignes_a_interroger(df, conf)
    if not len(concernees):
        log.info("LLM : aucune ligne à interroger")
        return df

    echecs = 0
    for index, ligne in concernees.iterrows():
        try:
            reponse = llm_client.extraire_equipement(ligne.to_dict(), conf)
        except llm_client.LlmError as exc:
            echecs += 1
            if echecs == 1:
                log.warning("LLM injoignable, colonnes %s laissées vides : %s",
                            config.LLM_PREFIXE + "*", exc)
            continue
        for champ, valeur in reponse.items():
            df.at[index, f"{config.LLM_PREFIXE}{champ}"] = valeur

    traitees = len(concernees) - echecs
    log.info("LLM : %d ligne(s) enrichie(s), %d échec(s)", traitees, echecs)
    return df


def load_template(path: Path | None = None) -> Workbook:
    """Ouvre le template en conservant les formules (data_only=False, jamais True)."""
    path = path or config.TEMPLATE_PATH
    if not path.exists():
        raise PipelineError(f"Template introuvable : {path}")
    wb = load_workbook(path, data_only=False, keep_vba=path.suffix.lower() == ".xlsm")
    log.info("Template chargé : %s (onglets : %s)", path.name, wb.sheetnames)
    return wb


def get_target_sheet(wb: Workbook, name: str | None = None) -> Worksheet:
    name = name if name is not None else config.TARGET_SHEET
    if name is None:
        return wb.worksheets[0]
    if name not in wb.sheetnames:
        raise PipelineError(f"Onglet '{name}' absent du template. Onglets : {wb.sheetnames}")
    return wb[name]


def is_formula_cell(cell) -> bool:
    """Vrai si la cellule contient une formule (classique, matricielle ou table de données)."""
    if isinstance(cell, MergedCell):
        return False
    v = cell.value
    if cell.data_type == "f":
        return True
    if isinstance(v, (ArrayFormula, DataTableFormula)):
        return True
    return isinstance(v, str) and v.startswith("=")


def to_excel_value(value: Any) -> Any:
    """Convertit une valeur pandas/numpy en type natif accepté par openpyxl."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if hasattr(value, "item"):          # numpy scalar -> python natif
        return value.item()
    return value


def write_value(ws: Worksheet, coord: str, value: Any) -> None:
    """Écrit dans une cellule d'input. Refuse toute cellule à formule ou fusionnée."""
    cell = ws[coord]
    if isinstance(cell, MergedCell):
        raise FormulaOverwriteError(
            f"{ws.title}!{coord} fait partie d'une plage fusionnée : "
            "écrire dans la cellule d'ancrage de la fusion."
        )
    if is_formula_cell(cell):
        raise FormulaOverwriteError(
            f"{ws.title}!{coord} contient une formule ({cell.value!r}) : écriture refusée."
        )
    cell.value = to_excel_value(value)


# ---------------------------------------------------------------------------
# 4. Remplissage
# ---------------------------------------------------------------------------
_AGGREGATORS = {
    "first": lambda s: s.iloc[0] if len(s) else None,
    "sum": lambda s: s.sum(),
    "max": lambda s: s.max(),
    "min": lambda s: s.min(),
    "count": lambda s: int(s.count()),
}


@dataclass
class RunResult:
    """Résumé d'un run, sérialisable en JSON pour un appelant externe (site web)."""
    status: str                      # ok | error | no_csv
    output_path: str | None = None
    output_name: str | None = None
    csv_name: str | None = None
    rows_written: int = 0
    cells_written: int = 0
    archived_to: str | None = None
    message: str | None = None

    def to_json(self) -> str:
        return json.dumps(asdict(self), ensure_ascii=False)


def fill_cells(ws: Worksheet, df: pd.DataFrame,
               mapping: dict | None = None) -> int:
    """Applique CELL_MAPPING : une colonne CSV -> une cellule fixe."""
    mapping = mapping if mapping is not None else config.CELL_MAPPING
    written = 0
    for column, spec in mapping.items():
        if isinstance(spec, str):
            coord, agg = spec, "first"
        else:
            coord, agg = spec["cell"], spec.get("agg", "first")
        if agg not in _AGGREGATORS:
            raise PipelineError(f"Agrégat inconnu '{agg}' pour la colonne {column}")
        value = _AGGREGATORS[agg](df[column])
        write_value(ws, coord, value)
        written += 1
    log.info("Cellules uniques écrites : %d", written)
    return written


def fill_table(ws: Worksheet, df: pd.DataFrame,
               mapping: dict | None = None) -> int:
    """Applique TABLE_MAPPING : chaque ligne du CSV -> une ligne Excel."""
    mapping = mapping if mapping is not None else config.TABLE_MAPPING
    start_row: int = mapping["start_row"]
    columns: dict[str, str] = mapping["columns"]
    max_rows: int | None = mapping.get("max_rows")

    if not columns:
        return 0
    if max_rows is not None and len(df) > max_rows:
        raise PipelineError(
            f"Le CSV contient {len(df)} lignes, le template n'en accepte que {max_rows} "
            f"(TABLE_MAPPING['max_rows'])."
        )

    if config.CLEAR_TABLE_BEFORE_WRITE and max_rows is not None:
        for r in range(start_row, start_row + max_rows):
            for col_letter in columns.values():
                write_value(ws, f"{col_letter}{r}", None)

    for i, (_, row) in enumerate(df.iterrows()):
        r = start_row + i
        for csv_col, col_letter in columns.items():
            write_value(ws, f"{col_letter}{r}", row[csv_col])
    log.info("Lignes de tableau écrites : %d (à partir de la ligne %d)", len(df), start_row)
    return len(df)


def fill_workbook(wb: Workbook, df: pd.DataFrame) -> tuple[int, int]:
    """Remplit l'onglet cible. Retourne (cellules uniques écrites, lignes écrites)."""
    ws = get_target_sheet(wb)
    n_cells = fill_cells(ws, df)
    n_rows = fill_table(ws, df)
    if config.FORCE_FULL_CALC_ON_LOAD:
        wb.calculation.fullCalcOnLoad = True
    return n_cells, n_rows


# ---------------------------------------------------------------------------
# 5. Sortie et archivage
# ---------------------------------------------------------------------------
def build_output_path(csv_path: Path, run_ts: datetime,
                      output_dir: Path | None = None,
                      template_path: Path | None = None) -> Path:
    output_dir = output_dir or config.OUTPUT_DIR
    template_path = template_path or config.TEMPLATE_PATH
    csv_date = extract_csv_date(csv_path) or run_ts.date()
    name = config.OUTPUT_NAME_PATTERN.format(
        template_stem=template_path.stem,
        csv_stem=csv_path.stem,
        date=csv_date.isoformat(),
        run_ts=run_ts.strftime("%Y%m%d_%H%M%S"),
    )
    return output_dir / name


def save_workbook(wb: Workbook, out_path: Path) -> Path:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    log.info("Fichier généré : %s", out_path)
    return out_path


def archive_csv(csv_path: Path, run_ts: datetime,
                archive_dir: Path | None = None) -> Path:
    """Déplace le CSV traité vers archive/YYYY-MM-DD/<horodatage>_<nom>."""
    archive_dir = archive_dir or config.ARCHIVE_DIR
    day_dir = archive_dir / run_ts.strftime("%Y-%m-%d")
    day_dir.mkdir(parents=True, exist_ok=True)
    dest = day_dir / f"{run_ts.strftime('%Y%m%d_%H%M%S')}_{csv_path.name}"
    shutil.move(str(csv_path), str(dest))
    log.info("CSV archivé : %s", dest)
    return dest


# ---------------------------------------------------------------------------
# 6. Orchestration
# ---------------------------------------------------------------------------
def process(csv_path: Path, *, dry_run: bool = False,
            run_ts: datetime | None = None,
            output_dir: Path | None = None) -> RunResult:
    """Traite un CSV donné de bout en bout. Retourne le résumé du run (status='ok').

    `dry_run` : ne pas archiver le CSV (usage manuel ou appel depuis le site web).
    `output_dir` : dossier de sortie, sinon config.OUTPUT_DIR.
    """
    run_ts = run_ts or datetime.now()
    df = load_source(csv_path)
    if not config.BLOCS.get("actif"):
        validate_columns(df)
    df = enrichir_avec_llm(df)
    wb = load_template()
    if config.BLOCS.get("actif"):
        n_cells, n_rows = fill_blocs(get_target_sheet(wb, config.BLOCS["feuille"]), df)
    else:
        n_cells, n_rows = fill_workbook(wb, df)
    out_path = build_output_path(csv_path, run_ts, output_dir=output_dir)
    save_workbook(wb, out_path)
    archived = None
    if dry_run:
        log.info("CSV non archivé (dry-run / no-archive) : %s", csv_path)
    else:
        archived = archive_csv(csv_path, run_ts)
    return RunResult(
        status="ok", output_path=str(out_path), output_name=out_path.name,
        csv_name=csv_path.name, rows_written=n_rows, cells_written=n_cells,
        archived_to=str(archived) if archived else None,
    )


def run(csv_override: Path | None = None, dry_run: bool = False,
        output_dir: Path | None = None, json_output: bool = False) -> int:
    """Point d'entrée : renvoie le code de sortie destiné au scheduler.

    En mode `json_output`, la dernière ligne de stdout est un JSON RunResult
    et les logs console partent sur stderr (pattern utilisé par nh-web).
    """
    setup_logging(console_stream=sys.stderr if json_output else None)
    started = datetime.now()
    log.info("=== RUN START ===")
    result: RunResult
    code: int
    try:
        csv_path = csv_override or find_latest_csv()
        if csv_path is None:
            log.warning("Aucun CSV trouvé dans %s (pattern %s). Run ignoré.",
                        config.INCOMING_DIR, config.CSV_GLOB)
            log.info("=== RUN SKIPPED | no_csv | %.2fs ===", (datetime.now() - started).total_seconds())
            result = RunResult(status="no_csv", message=f"Aucun CSV dans {config.INCOMING_DIR}")
            code = config.EXIT_CODE_NO_CSV
        else:
            if not csv_path.exists():
                raise PipelineError(f"CSV introuvable : {csv_path}")
            result = process(csv_path, dry_run=dry_run, run_ts=started, output_dir=output_dir)
            log.info("=== RUN OK | csv=%s | out=%s | %.2fs ===",
                     csv_path.name, result.output_name, (datetime.now() - started).total_seconds())
            code = 0
    except PipelineError as e:
        log.error("=== RUN FAILED | %s ===", e)
        result = RunResult(status="error", message=str(e),
                           csv_name=csv_override.name if csv_override else None)
        code = config.EXIT_CODE_ERROR
    except Exception as e:  # noqa: BLE001
        log.exception("=== RUN FAILED | erreur inattendue ===")
        result = RunResult(status="error", message=f"Erreur inattendue : {type(e).__name__}: {e}",
                           csv_name=csv_override.name if csv_override else None)
        code = config.EXIT_CODE_ERROR

    if json_output:
        print(result.to_json(), flush=True)
    return code


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Génère l'Excel du jour à partir du CSV entrant.")
    parser.add_argument("--csv", type=Path, default=None,
                        help="Traiter ce CSV précis au lieu de scanner le dossier entrant.")
    parser.add_argument("--dry-run", "--no-archive", dest="dry_run", action="store_true",
                        help="Génère l'Excel mais laisse le CSV en place (pas d'archivage).")
    parser.add_argument("--output-dir", type=Path, default=None,
                        help="Dossier de sortie (défaut : config.OUTPUT_DIR).")
    parser.add_argument("--json-output", action="store_true",
                        help="Imprime un résumé JSON en dernière ligne de stdout (appel depuis nh-web).")
    args = parser.parse_args(argv)
    return run(csv_override=args.csv, dry_run=args.dry_run,
               output_dir=args.output_dir, json_output=args.json_output)


if __name__ == "__main__":
    sys.exit(main())

# %% ----------------------------------------------------------------------
# Zone de test interactive (VS Code : exécuter cellule par cellule)
# ----------------------------------------------------------------------
# setup_logging()
# csv_path = find_latest_csv(); print(csv_path)
# df = load_csv(csv_path); df.head()
# validate_columns(df)
# wb = load_template(); ws = get_target_sheet(wb)
# fill_workbook(wb, df)
# save_workbook(wb, build_output_path(csv_path, datetime.now()))
