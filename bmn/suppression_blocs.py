# %% [markdown]
# # Blocs machine du template de dispo : suppression et duplication
#
# La dispo ne doit contenir que les machines présentes dans l'extraction : les
# autres blocs sont SUPPRIMÉS (pas masqués).
#
# openpyxl sait retirer des lignes (`delete_rows`) mais ne met à jour ni les
# formules, ni les fusions, ni les mises en forme conditionnelles, ni les
# hauteurs de ligne, ni la zone d'impression, et peut laisser des cellules des
# lignes supprimées. Ce module reproduit donc la suppression de lignes d'Excel :
#
# - les cellules des lignes conservées remontent ;
# - chaque référence de formule est recalculée ; une référence vers une ligne
#   supprimée devient `#REF!`, une plage partiellement supprimée rétrécit ;
# - les références vers une autre feuille ne bougent pas ;
# - fusions, mises en forme conditionnelles, hauteurs, commentaires et zone
#   d'impression suivent le même décalage.

# %%
from __future__ import annotations

import re

from openpyxl.cell.cell import MergedCell
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formula.tokenizer import Token, Tokenizer
from copy import copy

from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

REF = "#REF!"


# ---------------------------------------------------------------------------
# Choix des lignes à supprimer
# ---------------------------------------------------------------------------
def lignes_des_blocs_absents(ws: Worksheet, presentes: set[str], hauteur: int,
                             colonne_machine: str = "A",
                             motif: str = r"NH\s?\d+") -> set[int]:
    """Lignes des emplacements machine à supprimer.

    Les emplacements vont de la première machine à la fin de la dernière, par pas
    de `hauteur`. Un emplacement est supprimé si sa machine n'est pas dans
    `presentes`, ou s'il n'a pas de nom (ex. l'emplacement vide de NH05).
    """
    from daily_report import normaliser_machine  # import tardif : évite la boucle

    motif_re = re.compile(motif, re.IGNORECASE)
    bases = [ligne for ligne in range(1, ws.max_row + 1)
             if isinstance(ws[f"{colonne_machine}{ligne}"].value, str)
             and motif_re.fullmatch(ws[f"{colonne_machine}{ligne}"].value.strip())]
    if not bases:
        return set()

    presentes = {normaliser_machine(p) for p in presentes}
    a_supprimer: set[int] = set()
    for debut in range(bases[0], bases[-1] + hauteur, hauteur):
        nom = ws[f"{colonne_machine}{debut}"].value
        nom_ok = isinstance(nom, str) and motif_re.fullmatch(nom.strip())
        if not nom_ok or normaliser_machine(nom) not in presentes:
            a_supprimer.update(range(debut, debut + hauteur))
    return a_supprimer


# ---------------------------------------------------------------------------
# Table de décalage : ancienne ligne -> nouvelle ligne (None si supprimée)
# ---------------------------------------------------------------------------
def table_de_decalage(supprimees: set[int], max_row: int) -> dict[int, int | None]:
    table: dict[int, int | None] = {}
    retirees = 0
    for ligne in range(1, max_row + 1):
        if ligne in supprimees:
            table[ligne] = None
            retirees += 1
        else:
            table[ligne] = ligne - retirees
    return table


def _nouvelle(table: dict, ligne: int) -> int | None:
    # Au-delà de la dernière ligne connue : décalage du total supprimé.
    if ligne in table:
        return table[ligne]
    retirees = sum(1 for v in table.values() if v is None)
    return ligne - retirees


def _plage_lignes(table: dict, debut: int, fin: int) -> tuple[int, int] | None:
    """Plage de lignes après suppression (rétrécie), ou None si tout est supprimé."""
    gardees = [_nouvelle(table, l) for l in range(debut, fin + 1)]
    gardees = [l for l in gardees if l is not None]
    if not gardees:
        return None
    return min(gardees), max(gardees)


# ---------------------------------------------------------------------------
# Formules
# ---------------------------------------------------------------------------
_PARTIE = re.compile(r"^(\$?)([A-Za-z]{1,3})?(\$?)(\d+)?$")


def _decaler_reference(reference: str, table: dict, feuille: str) -> str:
    feuille_ref = None
    cible = reference
    if "!" in reference:
        feuille_ref, cible = reference.rsplit("!", 1)
        if feuille_ref.strip("'") != feuille:
            return reference                      # autre feuille : inchangé

    parties = cible.split(":")
    decoupees = [_PARTIE.match(p) for p in parties]
    if not all(decoupees) or not any(m.group(4) for m in decoupees):
        return reference                          # nom défini, colonne entière...

    prefixe = f"{feuille_ref}!" if feuille_ref else ""
    if len(parties) == 1:
        d1, col, d2, ligne = decoupees[0].groups()
        nouvelle = _nouvelle(table, int(ligne))
        return REF if nouvelle is None else f"{prefixe}{d1}{col or ''}{d2}{nouvelle}"

    (a1, ac, a2, al), (b1, bc, b2, bl) = (m.groups() for m in decoupees)
    if al is None or bl is None:
        return reference
    plage = _plage_lignes(table, int(al), int(bl))
    if plage is None:
        return REF
    return f"{prefixe}{a1}{ac or ''}{a2}{plage[0]}:{b1}{bc or ''}{b2}{plage[1]}"


def decaler_formule(formule: str, table: dict, feuille: str) -> str:
    """Réécrit les références d'une formule après suppression de lignes."""
    if not (isinstance(formule, str) and formule.startswith("=")):
        return formule
    jetons = Tokenizer(formule)
    for jeton in jetons.items:
        if jeton.type == Token.OPERAND and jeton.subtype == Token.RANGE:
            jeton.value = _decaler_reference(jeton.value, table, feuille)
    return jetons.render()


# ---------------------------------------------------------------------------
# Suppression sur la feuille
# ---------------------------------------------------------------------------
def supprimer_lignes(ws: Worksheet, supprimees: set[int]) -> None:
    """Supprime des lignes comme Excel (voir l'en-tête du module)."""
    if not supprimees:
        return
    max_row = ws.max_row
    table = table_de_decalage(set(supprimees), max_row)
    feuille = ws.title

    # 1. Fusions : on les retire le temps du déplacement, on les remettra décalées.
    fusions = [range_boundaries(str(r)) for r in ws.merged_cells.ranges]
    for plage in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(plage))

    # 2. Cellules : chaque cellule conservée remonte ; les autres disparaissent.
    anciennes = list(ws._cells.items())
    ws._cells = {}
    for (ligne, colonne), cellule in anciennes:
        nouvelle = _nouvelle(table, ligne)
        if nouvelle is None:
            continue
        cellule.row = nouvelle
        if isinstance(cellule.value, str) and cellule.value.startswith("="):
            cellule.value = decaler_formule(cellule.value, table, feuille)
        ws._cells[(nouvelle, colonne)] = cellule

    # 3. Fusions remises en place (rétrécies, ou abandonnées si supprimées).
    for min_col, min_row, max_col, max_row_f in fusions:
        plage = _plage_lignes(table, min_row, max_row_f)
        if plage is None:
            continue
        if plage[0] == plage[1] and min_col == max_col:
            continue
        ws.merge_cells(start_row=plage[0], start_column=min_col,
                       end_row=plage[1], end_column=max_col)

    # 4. Hauteurs et visibilité des lignes.
    dimensions = {r: d for r, d in ws.row_dimensions.items()}
    ws.row_dimensions.clear()
    for ligne, dim in dimensions.items():
        nouvelle = _nouvelle(table, ligne)
        if nouvelle is None:
            continue
        dim.index = nouvelle
        ws.row_dimensions[nouvelle] = dim

    # 5. Mises en forme conditionnelles : plages et formules décalées.
    anciennes_mfc = list(ws.conditional_formatting)
    ws.conditional_formatting = ConditionalFormattingList()
    for mfc in anciennes_mfc:
        plages = []
        for plage in str(mfc.sqref).split():
            min_col, min_row, max_col, max_row_p = range_boundaries(plage)
            lignes = _plage_lignes(table, min_row, max_row_p)
            if lignes is None:
                continue
            plages.append(f"{get_column_letter(min_col)}{lignes[0]}:"
                          f"{get_column_letter(max_col)}{lignes[1]}")
        if not plages:
            continue
        for regle in mfc.rules:
            if regle.formula:
                regle.formula = [decaler_formule(f"={f}", table, feuille)[1:]
                                 for f in regle.formula]
            ws.conditional_formatting.add(" ".join(plages), regle)

    # 6. Zone d'impression.
    if ws.print_area:
        zones = []
        for zone in str(ws.print_area).split(","):
            cible = zone.split("!")[-1].replace("$", "")
            min_col, min_row, max_col, max_row_z = range_boundaries(cible)
            lignes = _plage_lignes(table, min_row, max_row_z)
            if lignes is not None:
                zones.append(f"{get_column_letter(min_col)}{lignes[0]}:"
                             f"{get_column_letter(max_col)}{lignes[1]}")
        ws.print_area = zones or None


# ---------------------------------------------------------------------------
# Machines de l'extraction absentes du template
# ---------------------------------------------------------------------------
def machines_absentes(ws: Worksheet, presentes: set[str], hauteur: int,
                      colonne_machine: str = "A",
                      motif: str = r"NH\s?\d+") -> list[str]:
    """Machines de l'extraction qui n'ont pas d'emplacement, par ordre numérique."""
    from daily_report import normaliser_machine

    connues = {nom for nom, _ in _blocs_nommes(ws, colonne_machine, motif)}
    manquantes = {normaliser_machine(m) for m in presentes} - connues
    return sorted(manquantes, key=_cle_numerique)


def _cle_numerique(nom: str):
    chiffres = re.search(r"(\d+)", nom)
    return (int(chiffres.group(1)) if chiffres else 0, nom)


def _blocs_nommes(ws: Worksheet, colonne_machine: str, motif: str) -> list[tuple[str, int]]:
    """(nom normalisé, première ligne) de chaque emplacement portant un nom."""
    from daily_report import normaliser_machine

    motif_re = re.compile(motif, re.IGNORECASE)
    trouves = []
    for ligne in range(1, ws.max_row + 1):
        valeur = ws[f"{colonne_machine}{ligne}"].value
        if isinstance(valeur, str) and motif_re.fullmatch(valeur.strip()):
            trouves.append((normaliser_machine(valeur), ligne))
    return trouves


def place_du_bloc(ws: Worksheet, nom: str, hauteur: int,
                  colonne_machine: str = "A", motif: str = r"NH\s?\d+") -> int:
    """Ligne où insérer le bloc d'une machine, pour rester dans l'ordre numérique."""
    blocs = _blocs_nommes(ws, colonne_machine, motif)
    if not blocs:
        raise PipelineError("Aucun emplacement machine dans le template : duplication impossible.")
    cle = _cle_numerique(nom)
    suivants = [ligne for autre, ligne in blocs if _cle_numerique(autre) > cle]
    if suivants:
        return min(suivants)
    return max(ligne for _, ligne in blocs) + hauteur     # après le dernier bloc


# ---------------------------------------------------------------------------
# Insertion de lignes (mouvement inverse de la suppression)
# ---------------------------------------------------------------------------
def table_dinsertion(position: int, nombre: int, max_row: int) -> dict[int, int]:
    """Ancienne ligne -> nouvelle ligne après insertion de `nombre` lignes en `position`."""
    return {ligne: (ligne + nombre if ligne >= position else ligne)
            for ligne in range(1, max_row + 1)}


def inserer_lignes(ws: Worksheet, position: int, nombre: int) -> None:
    """Insère des lignes vides comme Excel : tout ce qui suit descend, formules comprises."""
    if nombre <= 0:
        return
    table = table_dinsertion(position, nombre, ws.max_row)
    feuille = ws.title

    fusions = [range_boundaries(str(r)) for r in ws.merged_cells.ranges]
    for plage in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(plage))

    # Cellules : de bas en haut, pour ne jamais écraser une ligne pas encore déplacée.
    for (ligne, colonne), cellule in sorted(ws._cells.items(), key=lambda kv: -kv[0][0]):
        nouvelle = table[ligne]
        if nouvelle == ligne:
            continue
        del ws._cells[(ligne, colonne)]
        cellule.row = nouvelle
        if isinstance(cellule.value, str) and cellule.value.startswith("="):
            cellule.value = decaler_formule(cellule.value, table, feuille)
        ws._cells[(nouvelle, colonne)] = cellule
    # Les formules restées au-dessus peuvent pointer plus bas : elles suivent aussi.
    for (ligne, colonne), cellule in ws._cells.items():
        if ligne < position and isinstance(cellule.value, str) and cellule.value.startswith("="):
            cellule.value = decaler_formule(cellule.value, table, feuille)

    for min_col, min_row, max_col, max_row_f in fusions:
        ws.merge_cells(start_row=table[min_row], start_column=min_col,
                       end_row=table[max_row_f], end_column=max_col)

    dimensions = {r: d for r, d in ws.row_dimensions.items()}
    ws.row_dimensions.clear()
    for ligne, dim in dimensions.items():
        nouvelle = table.get(ligne, ligne)
        dim.index = nouvelle
        ws.row_dimensions[nouvelle] = dim

    anciennes_mfc = list(ws.conditional_formatting)
    ws.conditional_formatting = ConditionalFormattingList()
    for mfc in anciennes_mfc:
        plages = []
        for plage in str(mfc.sqref).split():
            min_col, min_row, max_col, max_row_p = range_boundaries(plage)
            plages.append(f"{get_column_letter(min_col)}{table[min_row]}:"
                          f"{get_column_letter(max_col)}{table[max_row_p]}")
        for regle in mfc.rules:
            if regle.formula:
                regle.formula = [decaler_formule(f"={f}", table, feuille)[1:] for f in regle.formula]
            ws.conditional_formatting.add(" ".join(plages), regle)

    if ws.print_area:
        zones = []
        for zone in str(ws.print_area).split(","):
            cible = zone.split("!")[-1].replace("$", "")
            min_col, min_row, max_col, max_row_z = range_boundaries(cible)
            zones.append(f"{get_column_letter(min_col)}{table[min_row]}:"
                         f"{get_column_letter(max_col)}{table[max_row_z]}")
        ws.print_area = zones or None


# ---------------------------------------------------------------------------
# Duplication d'un bloc pour une machine nouvelle
# ---------------------------------------------------------------------------
def _decaler_formule_bloc(formule: str, modele: int, hauteur: int, delta: int,
                          feuille: str) -> str:
    """Décale les références INTERNES au bloc modèle ; les autres restent (en-tête, autres blocs)."""
    jetons = Tokenizer(formule)
    for jeton in jetons.items:
        if jeton.type != Token.OPERAND or jeton.subtype != Token.RANGE:
            continue
        if "!" in jeton.value:
            continue
        parties = jeton.value.split(":")
        decoupees = [_PARTIE.match(p) for p in parties]
        if not all(decoupees) or not all(m.group(4) for m in decoupees):
            continue
        lignes = [int(m.group(4)) for m in decoupees]
        if not all(modele <= l < modele + hauteur for l in lignes):
            continue
        morceaux = []
        for m in decoupees:
            d1, col, d2, ligne = m.groups()
            morceaux.append(f"{d1}{col or ''}{d2}{int(ligne) + delta}")
        jeton.value = ":".join(morceaux)
    return jetons.render()


def dupliquer_bloc(ws: Worksheet, nom: str, modele: int, hauteur: int, position: int,
                   colonne_machine: str = "A", colonnes_donnees=()) -> int:
    """Insère un bloc vierge copié du bloc `modele`, au nom de `nom`. Retourne sa première ligne.

    Copie la structure (intitulés, formules, styles, fusions, hauteurs, mises en
    forme conditionnelles) mais PAS les données saisies du modèle : le nouveau
    bloc arrive vide, prêt à être rempli par l'extraction.

    `colonnes_donnees` : colonnes où le modèle peut contenir des saisies (valeurs
    FIA, pièces et références des HIL/CIL). Elles sont vidées dans le bloc neuf,
    sauf les cellules à formule, qui ne sont jamais touchées.
    """
    inserer_lignes(ws, position, hauteur)
    if modele >= position:
        modele += hauteur                     # le modèle a été poussé par l'insertion
    delta = position - modele

    fusions_modele = [range_boundaries(str(r)) for r in ws.merged_cells.ranges
                      if modele <= range_boundaries(str(r))[1] < modele + hauteur]

    for decalage in range(hauteur):
        ligne_modele, ligne_cible = modele + decalage, position + decalage
        for colonne in range(1, ws.max_column + 1):
            source = ws.cell(row=ligne_modele, column=colonne)
            if source.value is None and not source.has_style:
                continue
            cible = ws.cell(row=ligne_cible, column=colonne)
            if isinstance(cible, MergedCell) or isinstance(source, MergedCell):
                continue          # une fusion se recopie par sa plage, pas cellule par cellule
            if source.has_style:
                cible._style = copy(source._style)
            valeur = source.value
            if isinstance(valeur, str) and valeur.startswith("="):
                cible.value = _decaler_formule_bloc(valeur, modele, hauteur, delta, ws.title)
            elif isinstance(valeur, str):
                cible.value = valeur          # intitulés (FH, ATT, HIL, CIL...)
            # les valeurs saisies (nombres, dates) ne sont pas recopiées
        dim = ws.row_dimensions.get(ligne_modele)
        if dim is not None and dim.height is not None:
            ws.row_dimensions[ligne_cible].height = dim.height

    for min_col, min_row, max_col, max_row_f in fusions_modele:
        ws.merge_cells(start_row=min_row + delta, start_column=min_col,
                       end_row=max_row_f + delta, end_column=max_col)

    for mfc in list(ws.conditional_formatting):
        plages = [p for p in str(mfc.sqref).split()
                  if modele <= range_boundaries(p)[1] < modele + hauteur]
        for plage in plages:
            min_col, min_row, max_col, max_row_p = range_boundaries(plage)
            nouvelle = (f"{get_column_letter(min_col)}{min_row + delta}:"
                        f"{get_column_letter(max_col)}{max_row_p + delta}")
            for regle in mfc.rules:
                ws.conditional_formatting.add(nouvelle, copy(regle))

    for colonne in colonnes_donnees:
        for decalage in range(hauteur):
            cellule = ws[f"{colonne}{position + decalage}"]
            if isinstance(cellule, MergedCell):
                continue          # lecture seule : la valeur vit dans la cellule d'ancrage
            if not (isinstance(cellule.value, str) and cellule.value.startswith("=")):
                cellule.value = None

    ws.cell(row=position, column=1).value = nom
    cellule_nom = ws[f"{colonne_machine}{position}"]
    if cellule_nom.value != nom:
        cellule_nom.value = nom
    return position


def ajuster_zone_impression(ws: Worksheet) -> None:
    """Étend la zone d'impression jusqu'à la dernière ligne utilisée.

    Après ajout d'un bloc à la fin de la feuille, la zone d'impression du
    template s'arrêterait avant : la nouvelle machine ne s'imprimerait pas.
    La zone n'est jamais réduite ici.
    """
    if not ws.print_area:
        return
    zones = []
    for zone in str(ws.print_area).split(","):
        cible = zone.split("!")[-1].replace("$", "")
        min_col, min_row, max_col, max_row_z = range_boundaries(cible)
        zones.append(f"{get_column_letter(min_col)}{min_row}:"
                     f"{get_column_letter(max_col)}{max(max_row_z, ws.max_row)}")
    ws.print_area = zones
