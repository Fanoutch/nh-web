"""Ouverture bornée des classeurs : seuls les fichiers de l'exécution, avec leurs limites.

Lecture seule, valeurs (jamais de formule évaluée), macros jamais lancées."""
from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from .registre import Contexte

LIMITE_OCTETS = 20 * 1024 * 1024
LIMITE_LIGNES = 100_000
EXTENSIONS = (".xlsx", ".xlsm", ".csv")


class ErreurClasseur(Exception):
    """Message déjà rédigé en français, présentable tel quel au modèle puis à l'utilisateur."""


@dataclass
class Classeur:
    nom: str
    feuilles: list[str]
    feuille: str
    entetes: list[str]
    lignes: list[list]


def resoudre(ctx: Contexte, nom: str, extensions: tuple[str, ...] = EXTENSIONS) -> Path:
    """Chemin d'un fichier de l'exécution (déposé ou produit), après tous les contrôles."""
    if not isinstance(nom, str) or not nom or Path(nom).name != nom:
        raise ErreurClasseur(f"« {nom} » n'est pas un nom de fichier valide : donnez le nom seul, sans dossier.")
    produit = nom in ctx.fichiers_produits
    if nom not in ctx.fichiers_autorises and not produit:
        dispo = ", ".join(ctx.fichiers_autorises + ctx.fichiers_produits) or "aucun"
        raise ErreurClasseur(f"Le fichier « {nom} » n'est pas disponible pour cette exécution. "
                             f"Fichiers disponibles : {dispo}.")
    base = (ctx.dossier_sortie if produit else ctx.dossier_execution)
    chemin = base / nom
    if not chemin.is_file():
        raise ErreurClasseur(f"Le fichier « {nom} » est introuvable.")
    if chemin.resolve().parent != base.resolve():
        raise ErreurClasseur(f"Le fichier « {nom} » est en dehors du dossier de l'exécution : il est refusé.")
    if chemin.suffix.lower() not in extensions:
        raise ErreurClasseur(f"Le format de « {nom} » n'est pas pris en charge par cet outil. "
                             f"Formats acceptés : {', '.join(extensions)}.")
    if chemin.stat().st_size > LIMITE_OCTETS:
        raise ErreurClasseur(f"Le fichier « {nom} » est trop volumineux "
                             f"(limite : {LIMITE_OCTETS // (1024 * 1024)} Mo).")
    return chemin


def _entetes_uniques(brut: list) -> list[str]:
    vus, resultat = {}, []
    for i, valeur in enumerate(brut, start=1):
        nom = str(valeur).strip() if valeur not in (None, "") else f"Colonne {i}"
        vus[nom] = vus.get(nom, 0) + 1
        resultat.append(nom if vus[nom] == 1 else f"{nom}_{vus[nom]}")
    return resultat


def _vide(ligne) -> bool:
    return all(c is None or (isinstance(c, str) and not c.strip()) for c in ligne)


def _mettre_en_forme(lignes_brutes, nom: str, feuilles: list[str], feuille: str) -> Classeur:
    entetes, lignes = None, []
    for brute in lignes_brutes:
        if _vide(brute):
            continue
        if entetes is None:
            entetes = _entetes_uniques(list(brute))
            continue
        if len(lignes) >= LIMITE_LIGNES:
            raise ErreurClasseur(f"Le fichier « {nom} » dépasse {LIMITE_LIGNES} lignes : "
                                 "il est trop grand pour être traité.")
        ligne = list(brute)[:len(entetes)]
        ligne += [None] * (len(entetes) - len(ligne))
        lignes.append(ligne)
    return Classeur(nom, feuilles, feuille, entetes or [], lignes)


def _lire_csv(chemin: Path, nom: str) -> Classeur:
    for encodage in ("utf-8-sig", "cp1252"):
        try:
            texte = chemin.read_text(encoding=encodage)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ErreurClasseur(f"Le fichier « {nom} » est illisible : son encodage n'est pas reconnu.")
    echantillon = texte[:4096]
    try:
        dialecte = csv.Sniffer().sniff(echantillon, delimiters=";,\t")
    except csv.Error:
        dialecte = csv.excel
        dialecte.delimiter = ";" if echantillon.count(";") >= echantillon.count(",") else ","
    lecteur = csv.reader(texte.splitlines(), dialecte)
    return _mettre_en_forme(lecteur, nom, [chemin.stem], chemin.stem)


def charger(ctx: Contexte, nom: str, feuille: str | None = None) -> Classeur:
    chemin = resoudre(ctx, nom)
    if chemin.suffix.lower() == ".csv":
        return _lire_csv(chemin, nom)
    try:
        wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True, keep_vba=False)
    except Exception as exc:  # noqa: BLE001 — format cassé, mot de passe, etc.
        raise ErreurClasseur(f"Le fichier « {nom} » est illisible. Vérifiez qu'il s'ouvre normalement dans Excel "
                             "et qu'il n'est pas protégé par un mot de passe.") from exc
    try:
        feuilles = list(wb.sheetnames)
        choisie = feuille or (feuilles[0] if feuilles else "")
        if choisie not in feuilles:
            raise ErreurClasseur(f"La feuille « {choisie} » n'existe pas dans « {nom} ». "
                                 f"Feuilles disponibles : {', '.join(feuilles)}.")
        return _mettre_en_forme(wb[choisie].iter_rows(values_only=True), nom, feuilles, choisie)
    finally:
        wb.close()
