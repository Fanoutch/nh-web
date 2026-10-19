"""Les huit outils Excel de l'agent.

Règle qui ne se discute pas : un classeur déposé n'est JAMAIS modifié. Tout outil qui « modifie » produit un
nouveau fichier dans le dossier de sortie de l'exécution. Le calcul reste du code ; le modèle ne tranche que
là où il y a du jugement (arbitrer des rapprochements douteux)."""
from __future__ import annotations

import functools
import re
import shutil
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl

from .classeur import Classeur, ErreurClasseur, charger, resoudre
from .condition import ErreurCondition, compiler as compiler_condition, trouver_colonne
from .formule import ErreurCalcul, ErreurFormule, compiler as compiler_formule
from .registre import Contexte, Reponse, erreur, ok, outil
from .texte import nombre, normaliser, pour_json

MAX_LIGNES_MODELE = 200      # lignes renvoyées au modèle par appel (le fichier, lui, peut en avoir 100 000)
MAX_GROUPES = 200
MAX_AFFICHES = 50            # éléments d'une liste de rapprochement montrés au modèle
SEUIL_DOUTEUX = 0.7


def _garde(fonction):
    """Transforme les erreurs « prévues » (fichier, condition, formule) en réponse d'erreur lisible."""
    @functools.wraps(fonction)
    def enveloppe(ctx, **kw):
        try:
            return fonction(ctx, **kw)
        except ErreurClasseur as exc:
            return erreur(str(exc), "Utilisez excel_decrire pour voir les fichiers et les feuilles disponibles.")
        except ErreurCondition as exc:
            return erreur(str(exc), "Corrigez la condition (colonne, opérateur, valeur) et relancez l'outil.")
        except ErreurFormule as exc:
            return erreur(str(exc), "Corrigez la formule et relancez l'outil.")
    return enveloppe


def _colonnes(c: Classeur, noms) -> list[int]:
    if not noms:
        return list(range(len(c.entetes)))
    if not isinstance(noms, list):
        raise ErreurCondition("« colonnes » doit être une liste de noms de colonnes.")
    return [trouver_colonne(n, c.entetes) for n in noms]


def _entier(valeur, defaut: int, mini: int, maxi: int | None = None) -> int:
    try:
        v = int(valeur)
    except (TypeError, ValueError):
        return defaut
    v = max(mini, v)
    return min(v, maxi) if maxi is not None else v


def _page(c: Classeur, lignes: list[list], colonnes, depuis, limite) -> Reponse:
    indices = _colonnes(c, colonnes)
    depuis = _entier(depuis, 0, 0)
    demande = _entier(limite, MAX_LIGNES_MODELE, 1)
    limite = min(demande, MAX_LIGNES_MODELE)
    tranche = lignes[depuis:depuis + limite]
    total = len(lignes)
    fin = depuis + len(tranche)
    suite = f"Lignes {depuis + 1} à {fin} sur {total}." if tranche else f"Aucune ligne à partir de {depuis} (total : {total})."
    if fin < total:
        suite += f" Pour la suite, utilisez depuis={fin}."
    if demande > MAX_LIGNES_MODELE:
        suite += f" (Limité à {MAX_LIGNES_MODELE} lignes par appel : filtrez avec excel_filtrer ou excel_selectionner.)"
    return ok(f"{len(tranche)} ligne(s) lues dans « {c.nom} » (feuille « {c.feuille} »).",
              {"colonnes": [c.entetes[i] for i in indices],
               "lignes": [[pour_json(l[i]) for i in indices] for l in tranche],
               "depuis": depuis, "retournees": len(tranche), "total": total}, suite)


# --------------------------------------------------------------------------- lecture

@outil("excel_decrire",
       "Sans argument : liste les classeurs disponibles. Avec « fichier » : feuilles, nombre de lignes, colonnes "
       "et 3 lignes d'aperçu de chaque feuille.",
       {"type": "object", "properties": {"fichier": {"type": "string", "description": "Nom du fichier à décrire"}}})
@_garde
def excel_decrire(ctx: Contexte, fichier: str | None = None) -> Reponse:
    if not fichier:
        liste = []
        for nom in ctx.fichiers_autorises + ctx.fichiers_produits:
            try:
                liste.append({"nom": nom, "taille_octets": resoudre(ctx, nom).stat().st_size,
                              "produit_par_l_agent": nom in ctx.fichiers_produits})
            except ErreurClasseur:
                continue
        return ok(f"{len(liste)} fichier(s) disponible(s).", {"fichiers": liste},
                  "Décrivez un fichier avec excel_decrire(fichier=...).")
    premier = charger(ctx, fichier)
    feuilles = []
    for nom_feuille in premier.feuilles:
        c = premier if nom_feuille == premier.feuille else charger(ctx, fichier, nom_feuille)
        feuilles.append({"nom": nom_feuille, "lignes": len(c.lignes), "colonnes": c.entetes,
                         "apercu": [dict(zip(c.entetes, [pour_json(v) for v in l])) for l in c.lignes[:3]]})
    return ok(f"« {fichier} » contient {len(feuilles)} feuille(s).", {"fichier": fichier, "feuilles": feuilles},
              "Lisez des lignes avec excel_lire, filtrez avec excel_filtrer ou excel_selectionner.")


@outil("excel_lire",
       f"Lit des lignes d'un classeur ({MAX_LIGNES_MODELE} au plus par appel). Paginez avec « depuis ».",
       {"type": "object", "properties": {
           "fichier": {"type": "string"}, "feuille": {"type": "string"},
           "colonnes": {"type": "array", "items": {"type": "string"}, "description": "Colonnes à garder"},
           "depuis": {"type": "integer", "description": "Nombre de lignes à sauter (0 au début)"},
           "limite": {"type": "integer"}}, "required": ["fichier"]})
@_garde
def excel_lire(ctx: Contexte, fichier: str, feuille: str | None = None, colonnes=None, depuis=0,
               limite=MAX_LIGNES_MODELE) -> Reponse:
    c = charger(ctx, fichier, feuille)
    return _page(c, c.lignes, colonnes, depuis, limite)


_CONDITION = {"type": "object", "description": "Condition : {colonne, operateur, valeur} ou "
              "{logique: et|ou, clauses: [...]}. Opérateurs : = != < <= > >= contient vide non_vide"}


@outil("excel_selectionner",
       "Renvoie les LIGNES qui vérifient une condition (le code filtre, vous ne voyez que le résultat).",
       {"type": "object", "properties": {
           "fichier": {"type": "string"}, "feuille": {"type": "string"}, "condition": _CONDITION,
           "colonnes": {"type": "array", "items": {"type": "string"}},
           "depuis": {"type": "integer"}, "limite": {"type": "integer"}}, "required": ["fichier"]})
@_garde
def excel_selectionner(ctx: Contexte, fichier: str, feuille: str | None = None, condition=None, colonnes=None,
                       depuis=0, limite=MAX_LIGNES_MODELE) -> Reponse:
    c = charger(ctx, fichier, feuille)
    test = compiler_condition(condition, c.entetes)
    return _page(c, [l for l in c.lignes if test(l)], colonnes, depuis, limite)


@outil("excel_filtrer",
       "Calcule sur les lignes qui vérifient une condition : compter, somme, moyenne, min, max, éventuellement par "
       "groupe. Ne renvoie JAMAIS les lignes (utilisez excel_selectionner pour cela).",
       {"type": "object", "properties": {
           "fichier": {"type": "string"}, "feuille": {"type": "string"}, "condition": _CONDITION,
           "calcul": {"type": "string", "enum": ["compter", "somme", "moyenne", "min", "max"]},
           "colonne": {"type": "string", "description": "Colonne numérique (inutile pour « compter »)"},
           "groupe": {"type": "string", "description": "Colonne de regroupement"}}, "required": ["fichier"]})
@_garde
def excel_filtrer(ctx: Contexte, fichier: str, feuille: str | None = None, condition=None, calcul="compter",
                  colonne: str | None = None, groupe: str | None = None) -> Reponse:
    if calcul not in ("compter", "somme", "moyenne", "min", "max"):
        return erreur(f"Le calcul « {calcul} » n'existe pas.", "Choisissez : compter, somme, moyenne, min ou max.")
    c = charger(ctx, fichier, feuille)
    test = compiler_condition(condition, c.entetes)
    retenues = [l for l in c.lignes if test(l)]
    ic = None
    if calcul != "compter":
        if not colonne:
            return erreur(f"Le calcul « {calcul} » demande une colonne numérique.", "Précisez l'argument « colonne ».")
        ic = trouver_colonne(colonne, c.entetes)
    ig = trouver_colonne(groupe, c.entetes) if groupe else None

    def calculer(lignes):
        if calcul == "compter":
            return len(lignes), 0
        valeurs = [nombre(l[ic]) for l in lignes]
        bonnes = [v for v in valeurs if v is not None]
        ignorees = len(valeurs) - len(bonnes)
        if not bonnes:
            return None, ignorees
        res = {"somme": sum(bonnes), "moyenne": sum(bonnes) / len(bonnes), "min": min(bonnes), "max": max(bonnes)}[calcul]
        return round(res, 6), ignorees

    if ig is None:
        valeur, ignorees = calculer(retenues)
        donnees = {"calcul": calcul, "colonne": colonne, "valeur": valeur, "lignes_retenues": len(retenues)}
        if ignorees:
            donnees["valeurs_non_numeriques_ignorees"] = ignorees
        return ok(f"{calcul} = {valeur} sur {len(retenues)} ligne(s) retenue(s).", donnees)
    groupes: dict[str, list] = {}
    for l in retenues:
        groupes.setdefault("(vide)" if l[ig] in (None, "") else str(l[ig]), []).append(l)
    if len(groupes) > MAX_GROUPES:
        return erreur(f"Il y a {len(groupes)} groupes différents : c'est trop pour les afficher (limite : {MAX_GROUPES}).",
                      "Choisissez une colonne de regroupement moins détaillée, ou ajoutez une condition.")
    return ok(f"{calcul} par « {groupe} » : {len(groupes)} groupe(s).",
              {"calcul": calcul, "colonne": colonne, "groupe": groupe,
               "groupes": {g: calculer(ls)[0] for g, ls in sorted(groupes.items())},
               "lignes_retenues": len(retenues)})


# --------------------------------------------------------------------------- écriture (fichiers neufs uniquement)

def _nom_sortie(ctx: Contexte, souhaite: str | None, defaut: str) -> str:
    base = Path(souhaite or defaut).name
    base = re.sub(r"[^\w.\- àâäéèêëîïôöùûüçÀÉÈÊ]", "_", base).strip(". ") or defaut
    base = Path(base).stem + ".xlsx"
    existants = set(ctx.fichiers_produits) | set(ctx.fichiers_autorises)
    nom, n = base, 1
    while nom in existants or (ctx.dossier_sortie / nom).exists():
        n += 1
        nom = f"{Path(base).stem}-{n}.xlsx"
    return nom


def _ecrire(ctx: Contexte, nom: str, entetes: list[str], lignes: list[list], feuille: str = "Feuille1") -> Path:
    ctx.dossier_sortie.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", "_", feuille)[:31] or "Feuille1"
    ws.append(entetes)
    for ligne in lignes:
        ws.append(ligne)
    for rangee in ws.iter_rows():
        for cellule in rangee:       # un texte qui commence par « = » reste du texte : jamais une formule
            if isinstance(cellule.value, str) and cellule.value.startswith("="):
                cellule.data_type = "s"
    chemin = ctx.dossier_sortie / nom
    wb.save(chemin)
    ctx.fichiers_produits.append(nom)
    return chemin


def _artefact(nom: str, intermediaire: bool, lignes: int) -> dict:
    return {"nom": nom, "type": "xlsx", "intermediaire": intermediaire, "lignes": lignes}


@outil("excel_colonne",
       "Crée un NOUVEAU classeur identique au fichier donné avec une colonne calculée en plus. La formule est "
       "arithmétique, les colonnes s'écrivent entre crochets : [Besoin] - [Stock]. Le calcul est fait par le code.",
       {"type": "object", "properties": {
           "fichier": {"type": "string"}, "feuille": {"type": "string"},
           "formule": {"type": "string", "description": "Ex. : max([Besoin] - [Stock], 0)"},
           "nom_colonne": {"type": "string"}, "nom_sortie": {"type": "string"}},
        "required": ["fichier", "formule", "nom_colonne"]})
@_garde
def excel_colonne(ctx: Contexte, fichier: str, formule: str, nom_colonne: str, feuille: str | None = None,
                  nom_sortie: str | None = None) -> Reponse:
    c = charger(ctx, fichier, feuille)
    calcul = compiler_formule(formule, c.entetes)
    if nom_colonne in c.entetes:
        return erreur(f"La colonne « {nom_colonne} » existe déjà.", "Choisissez un autre nom pour la colonne calculée.")
    resultats, erreurs, exemple = [], 0, ""
    for i, ligne in enumerate(c.lignes, start=1):
        try:
            v = calcul(ligne)
            resultats.append(int(v) if float(v).is_integer() else round(v, 6))
        except ErreurCalcul as exc:
            resultats.append(None)
            erreurs += 1
            exemple = exemple or f"ligne {i} : {exc}"
    if c.lignes and erreurs == len(c.lignes):
        return erreur(f"Le calcul a échoué sur toutes les lignes ({exemple}).",
                      "Vérifiez que les colonnes de la formule contiennent des nombres (excel_decrire).")
    nom = _nom_sortie(ctx, nom_sortie, f"{Path(fichier).stem}-calcule.xlsx")
    _ecrire(ctx, nom, c.entetes + [nom_colonne], [l + [r] for l, r in zip(c.lignes, resultats)], c.feuille)
    donnees = {"fichier_produit": nom, "colonne_ajoutee": nom_colonne, "lignes": len(c.lignes), "erreurs": erreurs}
    if erreurs:
        donnees["exemple_erreur"] = exemple
    return ok(f"Colonne « {nom_colonne} » calculée sur {len(c.lignes) - erreurs} ligne(s) dans « {nom} ».", donnees,
              f"Utilisez maintenant « {nom} » (excel_selectionner, excel_creer…). Le fichier d'origine est intact.",
              [_artefact(nom, True, len(c.lignes))])


@outil("excel_creer",
       "Crée un classeur NEUF avec les colonnes et les lignes données.",
       {"type": "object", "properties": {
           "nom": {"type": "string"}, "colonnes": {"type": "array", "items": {"type": "string"}},
           "lignes": {"type": "array", "items": {"type": "array", "items": {}}},
           "feuille": {"type": "string"}}, "required": ["nom", "colonnes", "lignes"]})
@_garde
def excel_creer(ctx: Contexte, nom: str, colonnes: list, lignes: list, feuille: str | None = None) -> Reponse:
    if not isinstance(colonnes, list) or not colonnes or not all(isinstance(x, str) and x.strip() for x in colonnes):
        return erreur("« colonnes » doit être une liste non vide de noms.", "Exemple : [\"Matériel\", \"Quantité\"].")
    if len(set(colonnes)) != len(colonnes):
        return erreur("Deux colonnes portent le même nom.", "Donnez un nom différent à chaque colonne.")
    if not isinstance(lignes, list) or not all(isinstance(l, list) for l in lignes):
        return erreur("« lignes » doit être une liste de lignes, chacune étant une liste de valeurs.", "")
    trop_longues = [i for i, l in enumerate(lignes, start=1) if len(l) > len(colonnes)]
    if trop_longues:
        return erreur(f"La ligne {trop_longues[0]} a plus de valeurs que de colonnes ({len(colonnes)}).",
                      "Corrigez les lignes et relancez l'outil.")
    pleines = [list(l) + [None] * (len(colonnes) - len(l)) for l in lignes]
    fichier = _nom_sortie(ctx, nom, "nouveau.xlsx")
    _ecrire(ctx, fichier, colonnes, pleines, feuille or "Feuille1")
    return ok(f"Classeur « {fichier} » créé ({len(pleines)} ligne(s)).", {"fichier_produit": fichier, "lignes": len(pleines)},
              "Le fichier est prêt à être téléchargé.", [_artefact(fichier, False, len(pleines))])


@outil("excel_regenerer",
       "Repart d'un classeur existant et en produit une VERSION MODIFIÉE dans un nouveau fichier. L'original n'est "
       "jamais touché. Modifications : {cellule: \"B3\", valeur: ..., feuille?} ou {ajouter_ligne: [...], feuille?}.",
       {"type": "object", "properties": {
           "fichier_source": {"type": "string"},
           "modifications": {"type": "array", "items": {"type": "object"}},
           "nom_sortie": {"type": "string"}}, "required": ["fichier_source", "modifications", "nom_sortie"]})
@_garde
def excel_regenerer(ctx: Contexte, fichier_source: str, modifications: list, nom_sortie: str) -> Reponse:
    source = resoudre(ctx, fichier_source)
    if not isinstance(modifications, list) or not modifications:
        return erreur("« modifications » doit être une liste non vide.",
                      "Exemple : [{\"cellule\": \"B3\", \"valeur\": 12}].")
    nom = _nom_sortie(ctx, nom_sortie, f"{source.stem}-modifie.xlsx")
    ctx.dossier_sortie.mkdir(parents=True, exist_ok=True)
    sortie = ctx.dossier_sortie / nom
    if source.suffix.lower() == ".csv":
        c = charger(ctx, fichier_source)
        _ecrire(ctx, nom, c.entetes, c.lignes, c.feuille)
        ctx.fichiers_produits.remove(nom)
    else:
        shutil.copyfile(source, sortie)            # la source n'est lue que par cette copie
    try:
        wb = openpyxl.load_workbook(sortie, keep_vba=False)
    except Exception:  # noqa: BLE001
        sortie.unlink(missing_ok=True)
        return erreur(f"Le fichier « {fichier_source} » ne peut pas être régénéré (illisible).",
                      "Vérifiez qu'il s'ouvre normalement dans Excel.")
    try:
        for i, m in enumerate(modifications, start=1):
            if not isinstance(m, dict):
                raise ValueError(f"La modification {i} doit être un objet.")
            nom_feuille = m.get("feuille") or wb.sheetnames[0]
            if nom_feuille not in wb.sheetnames:
                raise ValueError(f"La feuille « {nom_feuille} » n'existe pas (feuilles : {', '.join(wb.sheetnames)}).")
            ws = wb[nom_feuille]
            if "cellule" in m:
                if not re.fullmatch(r"[A-Za-z]{1,3}[1-9]\d{0,6}", str(m["cellule"])):
                    raise ValueError(f"« {m['cellule']} » n'est pas une cellule valide (exemple : B3).")
                ws[str(m["cellule"]).upper()] = m.get("valeur")
                if isinstance(m.get("valeur"), str) and m["valeur"].startswith("="):
                    ws[str(m["cellule"]).upper()].data_type = "s"
            elif "ajouter_ligne" in m and isinstance(m["ajouter_ligne"], list):
                ws.append(m["ajouter_ligne"])
            else:
                raise ValueError(f"La modification {i} n'est pas comprise : utilisez « cellule » + « valeur » "
                                 "ou « ajouter_ligne ».")
        wb.save(sortie)
    except ValueError as exc:
        sortie.unlink(missing_ok=True)
        return erreur(str(exc), "Corrigez la modification et relancez l'outil.")
    finally:
        wb.close()
    ctx.fichiers_produits.append(nom) if nom not in ctx.fichiers_produits else None
    return ok(f"Nouvelle version « {nom} » produite ({len(modifications)} modification(s)). L'original est intact.",
              {"fichier_produit": nom, "modifications": len(modifications)},
              "Vérifiez l'aperçu avant de proposer le téléchargement.", [_artefact(nom, False, 0)])


# --------------------------------------------------------------------------- rapprochement

_MOTS_VIDES = {"mm", "cm", "de", "du", "la", "le", "les", "des", "d", "l", "en"}


def _canonique(valeur) -> str:
    texte = normaliser(valeur)
    texte = re.sub(r"[^a-z0-9]+", " ", texte)
    texte = re.sub(r"(\d)\s*x\s*(\d)", r"\1x\2", texte)
    return " ".join(t for t in texte.split() if t not in _MOTS_VIDES)


def _score(a: str, b: str) -> float:
    ta, tb = set(a.split()), set(b.split())
    jaccard = len(ta & tb) / len(ta | tb) if ta | tb else 0.0
    return max(jaccard, SequenceMatcher(None, a, b).ratio())


@outil("excel_rapprocher",
       "Rapproche deux colonnes de deux fichiers dont les libellés sont écrits différemment. Renvoie "
       "{surs, douteux, absents}. Le code tranche l'évident ; à vous d'arbitrer les « douteux » (passez les paires "
       "validées dans « confirmes »). Avec « nom_sortie », produit un classeur joint (fichier A + colonnes de B).",
       {"type": "object", "properties": {
           "fichier_a": {"type": "string"}, "colonne_a": {"type": "string"}, "feuille_a": {"type": "string"},
           "fichier_b": {"type": "string"}, "colonne_b": {"type": "string"}, "feuille_b": {"type": "string"},
           "confirmes": {"type": "array", "items": {"type": "array", "items": {"type": "string"}},
                         "description": "Paires [valeur_a, valeur_b] douteuses que vous validez"},
           "nom_sortie": {"type": "string"}},
        "required": ["fichier_a", "colonne_a", "fichier_b", "colonne_b"]})
@_garde
def excel_rapprocher(ctx: Contexte, fichier_a: str, colonne_a: str, fichier_b: str, colonne_b: str,
                     feuille_a: str | None = None, feuille_b: str | None = None, confirmes=None,
                     nom_sortie: str | None = None) -> Reponse:
    a, b = charger(ctx, fichier_a, feuille_a), charger(ctx, fichier_b, feuille_b)
    ia, ib = trouver_colonne(colonne_a, a.entetes), trouver_colonne(colonne_b, b.entetes)
    ok_manuels = {(normaliser(p[0]), normaliser(p[1])) for p in (confirmes or []) if isinstance(p, list) and len(p) == 2}

    exact: dict[str, int] = {}          # libellé canonique → ligne de B (mêmes mots, même ordre)
    sans_ordre: dict[str, int] = {}     # mêmes mots dans un autre ordre
    par_mot: dict[str, set[str]] = {}
    for j, ligne in enumerate(b.lignes):
        canon = _canonique(ligne[ib])
        if not canon:
            continue
        exact.setdefault(canon, j)
        sans_ordre.setdefault(" ".join(sorted(canon.split())), j)
        for mot in canon.split():
            par_mot.setdefault(mot, set()).add(canon)

    appariements: list[tuple[str, int | None, float]] = []   # (statut, ligne de B, score)
    for ligne in a.lignes:
        canon = _canonique(ligne[ia])
        if not canon:
            appariements.append(("absent", None, 0.0))
            continue
        if canon in exact:
            appariements.append(("sûr", exact[canon], 1.0))
            continue
        tri = " ".join(sorted(canon.split()))
        if tri in sans_ordre:
            appariements.append(("sûr", sans_ordre[tri], 1.0))
            continue
        candidats = set().union(*(par_mot.get(m, set()) for m in canon.split())) if canon.split() else set()
        meilleur, score = None, 0.0
        for cand in candidats:
            s = _score(canon, cand)
            if s > score:
                meilleur, score = cand, s
        if meilleur is not None and score >= SEUIL_DOUTEUX:
            j = exact[meilleur]
            valide = (normaliser(ligne[ia]), normaliser(b.lignes[j][ib])) in ok_manuels
            appariements.append(("sûr" if valide else "douteux", j, round(score, 3)))
        else:
            appariements.append(("absent", None, 0.0))

    surs = [{"a": pour_json(l[ia]), "b": pour_json(b.lignes[j][ib])} for l, (s, j, _) in zip(a.lignes, appariements) if s == "sûr"]
    douteux = [{"a": pour_json(l[ia]), "b": pour_json(b.lignes[j][ib]), "score": sc}
               for l, (s, j, sc) in zip(a.lignes, appariements) if s == "douteux"]
    absents = [pour_json(l[ia]) for l, (s, _, _) in zip(a.lignes, appariements) if s == "absent"]
    donnees = {"nb_surs": len(surs), "nb_douteux": len(douteux), "nb_absents": len(absents),
               "surs": surs[:MAX_AFFICHES], "douteux": douteux[:MAX_AFFICHES], "absents": absents[:MAX_AFFICHES]}
    artefacts: list[dict] = []
    suite = ("Arbitrez les « douteux » : renvoyez les paires que vous validez dans « confirmes »."
             if douteux else "Aucun cas douteux à arbitrer.")
    if nom_sortie:
        nom = _nom_sortie(ctx, nom_sortie, "rapproche.xlsx")
        vides_b = [None] * len(b.entetes)
        entetes_b = [e if e not in a.entetes else f"{e} (B)" for e in b.entetes]
        lignes_jointes = []
        for l, (statut, j, _) in zip(a.lignes, appariements):
            proposition = b.lignes[j][ib] if statut == "douteux" else None
            lignes_jointes.append(l + [statut, proposition] + (b.lignes[j] if statut == "sûr" else vides_b))
        _ecrire(ctx, nom, a.entetes + ["Rapprochement", "Proposition"] + entetes_b, lignes_jointes, a.feuille)
        donnees["fichier_produit"] = nom
        artefacts = [_artefact(nom, True, len(lignes_jointes))]
        suite += f" Le classeur joint « {nom} » est disponible pour la suite (excel_colonne, excel_selectionner)."
    return ok(f"{len(surs)} rapprochement(s) sûr(s), {len(douteux)} douteux, {len(absents)} absent(s).", donnees, suite, artefacts)
