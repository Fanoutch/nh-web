"""Fabrique le corpus du banc : un contenu source → .docx, .txt et .pdf (via LibreOffice).

Usage (depuis assistant/) : .venv/bin/python -m banc.fabriquer_corpus
"""
from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK

DOSSIER = Path(__file__).resolve().parent
EN_TETE = "DOCUMENT FICTIF — BANC DE TEST"
SAUT = "[SAUT]"


def lire_source(texte: str) -> list[tuple[str, object]]:
    blocs: list[tuple[str, object]] = []
    tableau: list[list[str]] = []
    for ligne in texte.splitlines():
        ligne = ligne.rstrip()
        if ligne.startswith("|"):
            tableau.append([c.strip() for c in ligne.strip("|").split("|")])
            continue
        if tableau:
            blocs.append(("tableau", tableau))
            tableau = []
        if not ligne:
            continue
        if ligne == SAUT:
            blocs.append(("saut", None))
        elif ligne.startswith("# "):
            blocs.append(("titre", ligne[2:].strip()))
        else:
            morceaux: list[str | None] = []
            for i, partie in enumerate(ligne.split(SAUT)):
                if i:
                    morceaux.append(None)
                morceaux.append(partie.strip())
            blocs.append(("paragraphe", morceaux))
    if tableau:
        blocs.append(("tableau", tableau))
    return blocs


def ecrire_docx(blocs, chemin: Path) -> Path:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = EN_TETE
    for genre, contenu in blocs:
        if genre == "titre":
            doc.add_heading(contenu, level=1)
        elif genre == "saut":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif genre == "tableau":
            table = doc.add_table(rows=len(contenu), cols=len(contenu[0]))
            table.style = "Table Grid"
            for i, ligne in enumerate(contenu):
                for j, valeur in enumerate(ligne):
                    table.cell(i, j).text = valeur
        else:
            paragraphe = doc.add_paragraph()
            for partie in contenu:
                if partie is None:
                    paragraphe.add_run().add_break(WD_BREAK.PAGE)
                else:
                    paragraphe.add_run(partie + " ")
    chemin = Path(chemin)
    doc.save(str(chemin))
    return chemin


def ecrire_txt(blocs, chemin: Path) -> Path:
    lignes = [EN_TETE, ""]
    for genre, contenu in blocs:
        if genre == "titre":
            lignes += [contenu, ""]
        elif genre == "tableau":
            lignes += [" | ".join(l) for l in contenu] + [""]
        elif genre == "paragraphe":
            lignes += [" ".join(p for p in contenu if p is not None), ""]
    chemin = Path(chemin)
    chemin.write_text("\n".join(lignes), encoding="utf-8")
    return chemin


def convertir_pdf(docx: Path, dossier: Path) -> Path:
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(dossier), str(docx)],
                   check=True, capture_output=True, timeout=180)
    return dossier / (docx.stem + ".pdf")


def main() -> int:
    for dossier in ("pdf", "docx", "txt"):
        (DOSSIER / "corpus" / dossier).mkdir(parents=True, exist_ok=True)
    for source in sorted((DOSSIER / "sources").glob("*.txt")):
        blocs = lire_source(source.read_text(encoding="utf-8"))
        docx = ecrire_docx(blocs, DOSSIER / "corpus" / "docx" / f"{source.stem}.docx")
        ecrire_txt(blocs, DOSSIER / "corpus" / "txt" / f"{source.stem}.txt")
        pdf = convertir_pdf(docx, DOSSIER / "corpus" / "pdf")
        print(f"{source.stem} : docx, txt, {pdf.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
