import pytest

from rag import lecteurs
from tests.fabrique import docx_texte, pdf_texte, pdf_vide


def test_formats_acceptes():
    assert lecteurs.formats() == ["docx", "pdf", "txt"]


def test_lit_un_pdf_page_par_page(tmp_path):
    chemin = pdf_texte(tmp_path / "manuel.pdf", ["Le vol de maintenance", "est ordonne par le chef BMN"])

    pages = lecteurs.lire(chemin)

    assert len(pages) == 1
    assert pages[0]["page"] == 1
    assert "vol de maintenance" in pages[0]["texte"]


def test_pdf_sans_texte_est_refuse(tmp_path):
    with pytest.raises(lecteurs.LectureImpossible, match=r"Aucun texte trouvé dans ce PDF \(document scanné \?\)\."):
        lecteurs.lire(pdf_vide(tmp_path / "scan.pdf"))


def test_lit_un_word(tmp_path):
    chemin = docx_texte(tmp_path / "procedure.docx", ["Titre", "Le bordereau de visite est rempli par le technicien."])

    pages = lecteurs.lire(chemin)

    assert pages == [{"page": None, "texte": "Titre\n\nLe bordereau de visite est rempli par le technicien."}]


def test_lit_les_tableaux_word(tmp_path):
    from docx import Document
    document = Document()
    tableau = document.add_table(rows=1, cols=2)
    tableau.rows[0].cells[0].text = "Machine"
    tableau.rows[0].cells[1].text = "NH-12"
    document.save(str(tmp_path / "t.docx"))

    assert "Machine ; NH-12" in lecteurs.lire(tmp_path / "t.docx")[0]["texte"]


def test_lit_un_txt_avec_sauts_de_page(tmp_path):
    chemin = tmp_path / "note.txt"
    chemin.write_text("Page un\fPage deux", encoding="utf-8")

    assert lecteurs.lire(chemin) == [{"page": 1, "texte": "Page un"}, {"page": 2, "texte": "Page deux"}]


def test_ancien_word_refuse_avec_conseil(tmp_path):
    chemin = tmp_path / "vieux.doc"
    chemin.write_bytes(b"\xd0\xcf\x11\xe0")

    with pytest.raises(lecteurs.LectureImpossible, match=r"enregistrez-le en \.docx"):
        lecteurs.lire(chemin)


def test_format_inconnu_refuse(tmp_path):
    chemin = tmp_path / "image.png"
    chemin.write_bytes(b"x")

    with pytest.raises(lecteurs.LectureImpossible, match="non pris en charge"):
        lecteurs.lire(chemin)


def test_fichier_endommage(tmp_path):
    chemin = tmp_path / "casse.pdf"
    chemin.write_bytes(b"pas un pdf")

    with pytest.raises(lecteurs.LectureImpossible, match="illisible ou endommagé") as info:
        lecteurs.lire(chemin)
    assert info.value.__cause__ is not None


def test_fichier_absent(tmp_path):
    with pytest.raises(lecteurs.LectureImpossible, match="introuvable"):
        lecteurs.lire(tmp_path / "absent.pdf")


def test_lit_un_txt_windows_cp1252(tmp_path):
    chemin = tmp_path / "ancien.txt"
    chemin.write_bytes("Procédure : contrôle du bordereau".encode("cp1252"))

    assert lecteurs.lire(chemin) == [{"page": None, "texte": "Procédure : contrôle du bordereau"}]


from docx import Document

from rag.lecteurs import lignes_tableau, lire


def test_lignes_tableau():
    assert lignes_tableau([["Tâche", "Échéance", ""], ["Visite V1", "25 h", "x"], ["", "", ""]]) == [
        "Tâche : Visite V1 ; Échéance : 25 h ; Colonne 3 : x"]
    assert lignes_tableau([["seule", "", "ligne"]]) == ["seule ; ligne"]
    assert lignes_tableau([]) == []


def test_lignes_tableau_cle_valeur_a_deux_colonnes():
    assert lignes_tableau([["Pression", "12 bar"], ["D\u00e9bit", "30 l/min"]]) == [
        "Pression : 12 bar", "D\u00e9bit : 30 l/min"]
    assert lignes_tableau([["Pression", ""], ["", ""], ["Note", "x"]]) == ["Pression", "Note : x"]
    assert lignes_tableau([["seule", "ligne"]]) == ["seule ; ligne"]      # une seule ligne : inchang\u00e9


def test_lignes_tableau_ignore_le_separateur_markdown():
    assert lignes_tableau([["Code", "Libell\u00e9", "Action"], ["---", ":---:", " - "],
                           ["HYD-112", "Fuite", "D\u00e9pose"]]) == [
        "Code : HYD-112 ; Libell\u00e9 : Fuite ; Action : D\u00e9pose"]


def test_txt_tableau_markdown_sans_ligne_de_tirets(tmp_path):
    chemin = tmp_path / "m.txt"
    chemin.write_text("| Code | Libell\u00e9 | Action |\n|---|---|---|\n| ELE-031 | Alternateur | R\u00e9gulateur |\n",
                      encoding="utf-8")
    texte = lire(chemin)[0]["texte"]
    assert "Code : ELE-031 ; Libell\u00e9 : Alternateur ; Action : R\u00e9gulateur" in texte
    assert "---" not in texte


def test_docx_tableau_a_sa_place_avec_en_tetes(tmp_path):
    doc = Document()
    doc.add_paragraph("Avant le tableau.")
    t = doc.add_table(rows=3, cols=3)
    for i, ligne in enumerate([("Code", "Libellé", "Action"), ("HYD-112", "Fuite", "Dépose"),
                               ("ELE-031", "Alternateur", "Régulateur")]):
        for j, valeur in enumerate(ligne):
            t.cell(i, j).text = valeur
    k = doc.add_table(rows=2, cols=2)
    for i, (a, b) in enumerate([("Pression", "12 bar"), ("Débit", "30 l/min")]):
        k.cell(i, 0).text, k.cell(i, 1).text = a, b
    doc.add_paragraph("Après le tableau.")
    chemin = tmp_path / "t.docx"
    doc.save(str(chemin))

    texte = lire(chemin)[0]["texte"]
    assert texte.index("Avant") < texte.index("Code : HYD-112 ; Libellé : Fuite ; Action : Dépose") < texte.index("Après")
    assert "Code : ELE-031 ; Libellé : Alternateur ; Action : Régulateur" in texte
    assert texte.index("Pression : 12 bar") < texte.index("Débit : 30 l/min") < texte.index("Après")
    assert " | " not in texte


def test_docx_cellules_fusionnees_sans_plantage(tmp_path):
    doc = Document()
    t = doc.add_table(rows=2, cols=3)
    for j, entete in enumerate(["A", "B", "C"]):
        t.cell(0, j).text = entete
    t.cell(1, 1).merge(t.cell(1, 2)).text = "fusion"
    t.cell(1, 0).text = "x"
    chemin = tmp_path / "f.docx"
    doc.save(str(chemin))
    assert "A : x ; B : fusion" in lire(chemin)[0]["texte"]


def test_txt_bloc_de_tableau_reecrit_et_prose_intacte(tmp_path):
    chemin = tmp_path / "t.txt"
    chemin.write_text("Une phrase a | b | c isolée.\n\nCode | Libellé | Action\nHYD-112 | Fuite | Dépose\n"
                      "ELE-031 | Alternateur | Régulateur\n\nFin.", encoding="utf-8")
    texte = lire(chemin)[0]["texte"]
    assert "Une phrase a | b | c isolée." in texte
    assert "Code : HYD-112 ; Libellé : Fuite ; Action : Dépose" in texte
    assert "Code : ELE-031 ; Libellé : Alternateur ; Action : Régulateur" in texte
