import hashlib

import openpyxl
import pytest

import outils  # noqa: F401 — enregistre les outils
from outils.registre import Contexte, executer

TOUS = {"excel_decrire", "excel_lire", "excel_filtrer", "excel_selectionner", "excel_colonne",
        "excel_rapprocher", "excel_creer", "excel_regenerer"}


def classeur(chemin, lignes, feuille="Stock"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = feuille
    for l in lignes:
        ws.append(l)
    wb.save(chemin)


@pytest.fixture
def ctx(tmp_path):
    entree, sortie = tmp_path / "entree", tmp_path / "sortie"
    entree.mkdir()
    return Contexte(dossier_execution=entree, fichiers_autorises=[], outils_autorises=set(TOUS),
                    dossier_sortie=sortie)


def deposer(ctx, nom, lignes, feuille="Stock"):
    classeur(ctx.dossier_execution / nom, lignes, feuille)
    ctx.fichiers_autorises.append(nom)


def sha(chemin):
    return hashlib.sha256(chemin.read_bytes()).hexdigest()


def relire(chemin):
    ws = openpyxl.load_workbook(chemin).active
    return [list(r) for r in ws.iter_rows(values_only=True)]


STOCK = [["Matériel", "Quantité", "Besoin", "Service"], ["Joint torique 12x2", 3, 10, "A"], ["Vis M6", 12, 10, "A"],
         ["Rondelle 8", 0, 4, "B"], ["Écrou M6", 7, 10, "B"]]


# ----------------------------------------------------------------------------- décrire / lire

def test_decrire_sans_argument_liste_les_fichiers(ctx):
    deposer(ctx, "stock.xlsx", STOCK)
    r = executer("excel_decrire", {}, ctx)
    assert r.statut == "ok" and [f["nom"] for f in r.donnees["fichiers"]] == ["stock.xlsx"]


def test_decrire_un_fichier(ctx):
    deposer(ctx, "stock.xlsx", STOCK)
    r = executer("excel_decrire", {"fichier": "stock.xlsx"}, ctx)
    f = r.donnees["feuilles"][0]
    assert f["nom"] == "Stock" and f["lignes"] == 4 and f["colonnes"][0] == "Matériel"
    assert f["apercu"][0] == {"Matériel": "Joint torique 12x2", "Quantité": 3, "Besoin": 10, "Service": "A"}


def test_decrire_fichier_hors_liste(ctx):
    r = executer("excel_decrire", {"fichier": "secret.xlsx"}, ctx)
    assert r.statut == "erreur" and "secret.xlsx" in r.resume and "excel_decrire" in r.suite


def test_lire_plafonne_a_200_lignes_et_pagine(ctx):
    deposer(ctx, "g.xlsx", [["N"]] + [[i] for i in range(450)])
    r = executer("excel_lire", {"fichier": "g.xlsx", "limite": 5000}, ctx)
    assert r.donnees["retournees"] == 200 and r.donnees["total"] == 450
    assert "depuis=200" in r.suite and "200 lignes par appel" in r.suite
    r2 = executer("excel_lire", {"fichier": "g.xlsx", "depuis": 200}, ctx)
    assert r2.donnees["lignes"][0] == [200]
    r3 = executer("excel_lire", {"fichier": "g.xlsx", "depuis": 400}, ctx)
    assert r3.donnees["retournees"] == 50 and "depuis" not in r3.suite


def test_lire_colonnes_choisies_et_colonne_inconnue(ctx):
    deposer(ctx, "stock.xlsx", STOCK)
    r = executer("excel_lire", {"fichier": "stock.xlsx", "colonnes": ["quantite", "Matériel"]}, ctx)
    assert r.donnees["colonnes"] == ["Quantité", "Matériel"] and r.donnees["lignes"][0] == [3, "Joint torique 12x2"]
    r = executer("excel_lire", {"fichier": "stock.xlsx", "colonnes": ["Prix"]}, ctx)
    assert r.statut == "erreur" and "Matériel" in r.resume


def test_feuille_absente(ctx):
    deposer(ctx, "stock.xlsx", STOCK)
    r = executer("excel_lire", {"fichier": "stock.xlsx", "feuille": "Nope"}, ctx)
    assert r.statut == "erreur" and "Stock" in r.resume


# ----------------------------------------------------------------------------- filtrer / sélectionner

def test_selectionner_renvoie_des_lignes(ctx):
    deposer(ctx, "stock.xlsx", STOCK)
    r = executer("excel_selectionner", {"fichier": "stock.xlsx",
                 "condition": {"colonne": "Quantité", "operateur": "<", "valeur": 8}}, ctx)
    assert [l[0] for l in r.donnees["lignes"]] == ["Joint torique 12x2", "Rondelle 8", "Écrou M6"]


@pytest.mark.parametrize("calcul,colonne,attendu", [
    ("compter", None, 4), ("somme", "Quantité", 22), ("moyenne", "Quantité", 5.5), ("min", "Quantité", 0),
    ("max", "Besoin", 10)])
def test_filtrer_les_calculs(ctx, calcul, colonne, attendu):
    deposer(ctx, "stock.xlsx", STOCK)
    args = {"fichier": "stock.xlsx", "calcul": calcul}
    if colonne:
        args["colonne"] = colonne
    r = executer("excel_filtrer", args, ctx)
    assert r.donnees["valeur"] == attendu


def test_filtrer_ne_renvoie_jamais_de_lignes(ctx):
    deposer(ctx, "stock.xlsx", STOCK)
    r = executer("excel_filtrer", {"fichier": "stock.xlsx"}, ctx)
    assert "lignes" not in r.donnees and "apercu" not in r.donnees


def test_filtrer_par_groupe_et_condition(ctx):
    deposer(ctx, "stock.xlsx", STOCK)
    r = executer("excel_filtrer", {"fichier": "stock.xlsx", "calcul": "somme", "colonne": "Besoin", "groupe": "Service",
                 "condition": {"colonne": "Quantité", "operateur": "<", "valeur": 8}}, ctx)
    assert r.donnees["groupes"] == {"A": 10, "B": 14}


def test_filtrer_erreurs_lisibles(ctx):
    deposer(ctx, "stock.xlsx", STOCK)
    assert "colonne" in executer("excel_filtrer", {"fichier": "stock.xlsx", "calcul": "somme"}, ctx).resume
    assert "n'existe pas" in executer("excel_filtrer", {"fichier": "stock.xlsx", "calcul": "médiane"}, ctx).resume
    r = executer("excel_filtrer", {"fichier": "stock.xlsx", "condition": {"colonne": "Prix", "operateur": "=", "valeur": 1}}, ctx)
    assert r.statut == "erreur" and "Prix" in r.resume


def test_filtrer_trop_de_groupes(ctx):
    deposer(ctx, "g.xlsx", [["N"]] + [[i] for i in range(300)])
    r = executer("excel_filtrer", {"fichier": "g.xlsx", "groupe": "N"}, ctx)
    assert r.statut == "erreur" and "300" in r.resume


# ----------------------------------------------------------------------------- colonne calculée

def test_colonne_calculee_en_nouveau_fichier_original_intact(ctx):
    deposer(ctx, "stock.xlsx", STOCK)
    avant = sha(ctx.dossier_execution / "stock.xlsx")
    r = executer("excel_colonne", {"fichier": "stock.xlsx", "formule": "max([Besoin] - [Quantité], 0)",
                                   "nom_colonne": "Manquant"}, ctx)
    assert r.statut == "ok" and r.donnees["erreurs"] == 0
    assert relire(ctx.dossier_sortie / r.donnees["fichier_produit"])[1:] == [
        ["Joint torique 12x2", 3, 10, "A", 7], ["Vis M6", 12, 10, "A", 0], ["Rondelle 8", 0, 4, "B", 4],
        ["Écrou M6", 7, 10, "B", 3]]
    assert sha(ctx.dossier_execution / "stock.xlsx") == avant
    # le fichier produit est utilisable par la suite
    r2 = executer("excel_selectionner", {"fichier": r.donnees["fichier_produit"],
                  "condition": {"colonne": "Manquant", "operateur": ">", "valeur": 0}}, ctx)
    assert len(r2.donnees["lignes"]) == 3


def test_colonne_erreurs_par_ligne_et_echec_total(ctx):
    deposer(ctx, "s.xlsx", [["Matériel", "Q", "B"], ["a", 1, 5], ["b", "beaucoup", 5]])
    r = executer("excel_colonne", {"fichier": "s.xlsx", "formule": "[B] - [Q]", "nom_colonne": "M"}, ctx)
    assert r.donnees["erreurs"] == 1 and "ligne 2" in r.donnees["exemple_erreur"]
    deposer(ctx, "t.xlsx", [["Matériel", "Q", "B"], ["a", "x", 5]])
    r = executer("excel_colonne", {"fichier": "t.xlsx", "formule": "[B] - [Q]", "nom_colonne": "M"}, ctx)
    assert r.statut == "erreur" and "toutes les lignes" in r.resume


def test_colonne_formule_dangereuse_ou_nom_deja_pris(ctx):
    deposer(ctx, "stock.xlsx", STOCK)
    r = executer("excel_colonne", {"fichier": "stock.xlsx", "formule": "__import__('os')", "nom_colonne": "X"}, ctx)
    assert r.statut == "erreur" and "non autorisé" in r.resume
    r = executer("excel_colonne", {"fichier": "stock.xlsx", "formule": "[Besoin]", "nom_colonne": "Besoin"}, ctx)
    assert r.statut == "erreur" and "existe déjà" in r.resume


# ----------------------------------------------------------------------------- créer

def test_creer_un_classeur_neuf(ctx):
    r = executer("excel_creer", {"nom": "commande", "colonnes": ["Matériel", "Quantité"],
                                 "lignes": [["Joint", 7], ["Vis"]]}, ctx)
    assert r.statut == "ok" and r.donnees["fichier_produit"] == "commande.xlsx"
    assert relire(ctx.dossier_sortie / "commande.xlsx") == [["Matériel", "Quantité"], ["Joint", 7], ["Vis", None]]
    assert r.artefacts[0]["nom"] == "commande.xlsx" and r.artefacts[0]["intermediaire"] is False


def test_creer_ne_depasse_jamais_le_dossier_et_n_ecrase_pas(ctx):
    r1 = executer("excel_creer", {"nom": "../../evil", "colonnes": ["A"], "lignes": [[1]]}, ctx)
    assert r1.donnees["fichier_produit"] == "evil.xlsx" and (ctx.dossier_sortie / "evil.xlsx").exists()
    r2 = executer("excel_creer", {"nom": "evil", "colonnes": ["A"], "lignes": [[2]]}, ctx)
    assert r2.donnees["fichier_produit"] == "evil-2.xlsx"


def test_creer_un_texte_commencant_par_egal_reste_du_texte(ctx):
    executer("excel_creer", {"nom": "f", "colonnes": ["A"], "lignes": [["=1+1"], ["=cmd|' /C calc'!A0"]]}, ctx)
    ws = openpyxl.load_workbook(ctx.dossier_sortie / "f.xlsx").active
    assert ws["A2"].data_type == "s" and ws["A3"].data_type == "s"


@pytest.mark.parametrize("args,mot", [
    ({"nom": "x", "colonnes": [], "lignes": []}, "colonnes"),
    ({"nom": "x", "colonnes": ["A", "A"], "lignes": []}, "même nom"),
    ({"nom": "x", "colonnes": ["A"], "lignes": [[1, 2]]}, "ligne 1"),
    ({"nom": "x", "colonnes": ["A"], "lignes": "oups"}, "lignes"),
])
def test_creer_entrees_invalides(ctx, args, mot):
    r = executer("excel_creer", args, ctx)
    assert r.statut == "erreur" and mot in r.resume


# ----------------------------------------------------------------------------- régénérer

def test_regenerer_original_intact_et_modifications_appliquees(ctx):
    deposer(ctx, "stock.xlsx", STOCK)
    avant = sha(ctx.dossier_execution / "stock.xlsx")
    r = executer("excel_regenerer", {"fichier_source": "stock.xlsx", "nom_sortie": "stock-v2",
                 "modifications": [{"cellule": "B2", "valeur": 99}, {"ajouter_ligne": ["Neuf", 1, 2, "C"]}]}, ctx)
    assert r.statut == "ok" and "intact" in r.resume
    sortie = relire(ctx.dossier_sortie / "stock-v2.xlsx")
    assert sortie[1][1] == 99 and sortie[-1] == ["Neuf", 1, 2, "C"]
    assert sha(ctx.dossier_execution / "stock.xlsx") == avant
    assert relire(ctx.dossier_execution / "stock.xlsx")[1][1] == 3


def test_regenerer_source_csv(ctx):
    (ctx.dossier_execution / "c.csv").write_text("A;B\n1;2\n", encoding="utf-8")
    ctx.fichiers_autorises.append("c.csv")
    r = executer("excel_regenerer", {"fichier_source": "c.csv", "nom_sortie": "c2",
                 "modifications": [{"cellule": "B2", "valeur": "9"}]}, ctx)
    assert r.statut == "ok" and relire(ctx.dossier_sortie / "c2.xlsx") == [["A", "B"], ["1", "9"]]
    assert (ctx.dossier_execution / "c.csv").read_text(encoding="utf-8") == "A;B\n1;2\n"


@pytest.mark.parametrize("modif,mot", [
    ({"cellule": "ZZ", "valeur": 1}, "cellule valide"),
    ({"cellule": "A1", "valeur": 1, "feuille": "Nope"}, "Nope"),
    ({"truc": 1}, "pas comprise"),
])
def test_regenerer_modification_invalide_ne_laisse_rien(ctx, modif, mot):
    deposer(ctx, "stock.xlsx", STOCK)
    r = executer("excel_regenerer", {"fichier_source": "stock.xlsx", "nom_sortie": "v2", "modifications": [modif]}, ctx)
    assert r.statut == "erreur" and mot in r.resume
    assert not (ctx.dossier_sortie / "v2.xlsx").exists() and ctx.fichiers_produits == []


def test_regenerer_source_hors_liste_refusee(ctx):
    r = executer("excel_regenerer", {"fichier_source": "autre.xlsx", "nom_sortie": "v", "modifications": [{"cellule": "A1", "valeur": 1}]}, ctx)
    assert r.statut == "erreur"


def test_un_nom_de_sortie_ne_peut_pas_recouvrir_un_fichier_depose(ctx):
    deposer(ctx, "stock.xlsx", STOCK)
    r = executer("excel_regenerer", {"fichier_source": "stock.xlsx", "nom_sortie": "stock",
                 "modifications": [{"cellule": "B2", "valeur": 1}]}, ctx)
    assert r.donnees["fichier_produit"] == "stock-2.xlsx"


# ----------------------------------------------------------------------------- rapprocher

STOCK2 = [["Matériel", "Quantité"], ["JOINT TORIQUE 12X2", 3], ["Vis M6 x 20", 12], ["Rondelle plate 8", 0],
          ["Écrou frein M6", 7], ["Graisse", 1]]
CATALOGUE = [["Désignation", "Réf"], ["Joint torique 12x2 mm", "R1"], ["Vis M6x20", "R2"],
             ["Rondelle 8 plate", "R3"], ["Écrou M6", "R4"], ["Tournevis", "R5"]]


def test_rapprocher_surs_douteux_absents(ctx):
    deposer(ctx, "stock.xlsx", STOCK2)
    deposer(ctx, "cat.xlsx", CATALOGUE, "Cat")
    r = executer("excel_rapprocher", {"fichier_a": "stock.xlsx", "colonne_a": "Matériel",
                 "fichier_b": "cat.xlsx", "colonne_b": "Désignation"}, ctx)
    d = r.donnees
    surs = {p["a"]: p["b"] for p in d["surs"]}
    assert surs["JOINT TORIQUE 12X2"] == "Joint torique 12x2 mm" and surs["Vis M6 x 20"] == "Vis M6x20"
    assert surs["Rondelle plate 8"] == "Rondelle 8 plate"
    assert [p["a"] for p in d["douteux"]] == ["Écrou frein M6"] and d["douteux"][0]["b"] == "Écrou M6"
    assert d["absents"] == ["Graisse"]
    assert "Arbitrez" in r.suite


def test_rapprocher_confirmes_et_classeur_joint(ctx):
    deposer(ctx, "stock.xlsx", STOCK2)
    deposer(ctx, "cat.xlsx", CATALOGUE, "Cat")
    avant = sha(ctx.dossier_execution / "stock.xlsx")
    r = executer("excel_rapprocher", {"fichier_a": "stock.xlsx", "colonne_a": "Matériel", "fichier_b": "cat.xlsx",
                 "colonne_b": "Désignation", "confirmes": [["Écrou frein M6", "Écrou M6"]], "nom_sortie": "joint"}, ctx)
    assert r.donnees["nb_douteux"] == 0 and r.donnees["nb_surs"] == 4
    joint = relire(ctx.dossier_sortie / "joint.xlsx")
    assert joint[0] == ["Matériel", "Quantité", "Rapprochement", "Proposition", "Désignation", "Réf"]
    par_nom = {l[0]: l for l in joint[1:]}
    assert par_nom["Écrou frein M6"][2:] == ["sûr", None, "Écrou M6", "R4"]
    assert par_nom["Graisse"][2:] == ["absent", None, None, None]
    assert sha(ctx.dossier_execution / "stock.xlsx") == avant


def test_rapprocher_douteux_non_confirme_n_est_pas_utilise(ctx):
    deposer(ctx, "stock.xlsx", STOCK2)
    deposer(ctx, "cat.xlsx", CATALOGUE, "Cat")
    executer("excel_rapprocher", {"fichier_a": "stock.xlsx", "colonne_a": "Matériel", "fichier_b": "cat.xlsx",
             "colonne_b": "Désignation", "nom_sortie": "joint"}, ctx)
    ligne = {l[0]: l for l in relire(ctx.dossier_sortie / "joint.xlsx")[1:]}["Écrou frein M6"]
    assert ligne[2:] == ["douteux", "Écrou M6", None, None]


def test_rapprocher_colonne_inconnue(ctx):
    deposer(ctx, "stock.xlsx", STOCK2)
    deposer(ctx, "cat.xlsx", CATALOGUE, "Cat")
    r = executer("excel_rapprocher", {"fichier_a": "stock.xlsx", "colonne_a": "Prix", "fichier_b": "cat.xlsx",
                 "colonne_b": "Désignation"}, ctx)
    assert r.statut == "erreur" and "Prix" in r.resume


# ----------------------------------------------------------------------------- cas de référence, sans LLM

def test_cas_stock_vers_commande_de_bout_en_bout(ctx):
    """Ce que l'agent enchaînera : rapprocher → calculer le manquant → ne garder que les manques → bon de commande."""
    deposer(ctx, "stock.xlsx", [["Matériel", "En rayon", "Besoin"], ["JOINT TORIQUE 12X2", 3, 10], ["Vis M6 x 20", 12, 10],
                                ["Rondelle plate 8", 0, 4], ["Écrou frein M6", 7, 10]])
    deposer(ctx, "catalogue.xlsx", CATALOGUE, "Cat")
    originaux = {n: sha(ctx.dossier_execution / n) for n in ctx.fichiers_autorises}

    r = executer("excel_rapprocher", {"fichier_a": "stock.xlsx", "colonne_a": "Matériel", "fichier_b": "catalogue.xlsx",
                 "colonne_b": "Désignation", "confirmes": [["Écrou frein M6", "Écrou M6"]], "nom_sortie": "joint"}, ctx)
    assert r.statut == "ok"
    r = executer("excel_colonne", {"fichier": "joint.xlsx", "formule": "max([Besoin] - [En rayon], 0)",
                 "nom_colonne": "À commander"}, ctx)
    assert r.statut == "ok"
    r = executer("excel_selectionner", {"fichier": r.donnees["fichier_produit"],
                 "condition": {"colonne": "À commander", "operateur": ">", "valeur": 0},
                 "colonnes": ["Réf", "Désignation", "À commander"]}, ctx)
    assert r.donnees["lignes"] == [["R1", "Joint torique 12x2 mm", 7], ["R3", "Rondelle 8 plate", 4],
                                   ["R4", "Écrou M6", 3]]
    r = executer("excel_creer", {"nom": "bon-de-commande", "colonnes": r.donnees["colonnes"], "lignes": r.donnees["lignes"]}, ctx)
    assert relire(ctx.dossier_sortie / "bon-de-commande.xlsx") == [
        ["Réf", "Désignation", "À commander"], ["R1", "Joint torique 12x2 mm", 7], ["R3", "Rondelle 8 plate", 4],
        ["R4", "Écrou M6", 3]]
    assert {n: sha(ctx.dossier_execution / n) for n in ctx.fichiers_autorises} == originaux   # jamais modifiés
