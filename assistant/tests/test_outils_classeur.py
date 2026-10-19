import os

import openpyxl
import pytest

from outils import classeur
from outils.classeur import ErreurClasseur, charger, resoudre
from outils.registre import Contexte


def faire_ctx(tmp_path, noms):
    return Contexte(dossier_execution=tmp_path, fichiers_autorises=list(noms), outils_autorises=set(),
                    dossier_sortie=tmp_path / "sortie")


def ecrire_xlsx(chemin, lignes, feuille="Stock"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = feuille
    for ligne in lignes:
        ws.append(ligne)
    wb.save(chemin)


def test_lecture_xlsx_en_tetes_et_lignes(tmp_path):
    ecrire_xlsx(tmp_path / "s.xlsx", [["Matériel", "Quantité"], ["Joint", 3], ["Vis", 10]])
    c = charger(faire_ctx(tmp_path, ["s.xlsx"]), "s.xlsx")
    assert c.entetes == ["Matériel", "Quantité"]
    assert c.lignes == [["Joint", 3], ["Vis", 10]]
    assert c.feuille == "Stock" and c.feuilles == ["Stock"]


def test_en_tetes_vides_et_doublons_sont_rendus_uniques(tmp_path):
    ecrire_xlsx(tmp_path / "s.xlsx", [["A", "A", None], [1, 2, 3]])
    c = charger(faire_ctx(tmp_path, ["s.xlsx"]), "s.xlsx")
    assert c.entetes == ["A", "A_2", "Colonne 3"]


def test_lignes_vides_ignorees_et_lignes_courtes_completees(tmp_path):
    ecrire_xlsx(tmp_path / "s.xlsx", [[None, None], ["A", "B"], [1], [None, None], [2, 3]])
    c = charger(faire_ctx(tmp_path, ["s.xlsx"]), "s.xlsx")
    assert c.entetes == ["A", "B"] and c.lignes == [[1, None], [2, 3]]


def test_classeur_vide(tmp_path):
    ecrire_xlsx(tmp_path / "v.xlsx", [])
    c = charger(faire_ctx(tmp_path, ["v.xlsx"]), "v.xlsx")
    assert c.entetes == [] and c.lignes == []


def test_feuille_absente_liste_les_feuilles(tmp_path):
    ecrire_xlsx(tmp_path / "s.xlsx", [["A"], [1]])
    with pytest.raises(ErreurClasseur) as exc:
        charger(faire_ctx(tmp_path, ["s.xlsx"]), "s.xlsx", "Autre")
    assert "Autre" in str(exc.value) and "Stock" in str(exc.value)


def test_lecture_csv_point_virgule(tmp_path):
    (tmp_path / "c.csv").write_text("Matériel;Quantité\nJoint;3\nVis;10\n", encoding="utf-8")
    c = charger(faire_ctx(tmp_path, ["c.csv"]), "c.csv")
    assert c.entetes == ["Matériel", "Quantité"] and c.lignes == [["Joint", "3"], ["Vis", "10"]]


def test_fichier_hors_liste_refuse(tmp_path):
    ecrire_xlsx(tmp_path / "s.xlsx", [["A"], [1]])
    with pytest.raises(ErreurClasseur) as exc:
        resoudre(faire_ctx(tmp_path, ["autre.xlsx"]), "s.xlsx")
    assert "autre.xlsx" in str(exc.value)


@pytest.mark.parametrize("nom", ["../s.xlsx", "/etc/passwd", "sous/s.xlsx"])
def test_chemin_refuse(tmp_path, nom):
    with pytest.raises(ErreurClasseur):
        resoudre(faire_ctx(tmp_path, [nom]), nom)


def test_lien_symbolique_hors_dossier_refuse(tmp_path):
    dossier = tmp_path / "exec"
    dossier.mkdir()
    secret = tmp_path / "secret.xlsx"
    ecrire_xlsx(secret, [["A"], [1]])
    os.symlink(secret, dossier / "lien.xlsx")
    with pytest.raises(ErreurClasseur):
        resoudre(faire_ctx(dossier, ["lien.xlsx"]), "lien.xlsx")


def test_extension_refusee(tmp_path):
    (tmp_path / "x.exe").write_bytes(b"MZ")
    with pytest.raises(ErreurClasseur) as exc:
        resoudre(faire_ctx(tmp_path, ["x.exe"]), "x.exe")
    assert ".xlsx" in str(exc.value)


def test_fichier_trop_gros_refuse_avant_lecture(tmp_path, monkeypatch):
    monkeypatch.setattr(classeur, "LIMITE_OCTETS", 100)
    (tmp_path / "g.csv").write_bytes(b"a;b\n" + b"x" * 101)
    with pytest.raises(ErreurClasseur) as exc:
        resoudre(faire_ctx(tmp_path, ["g.csv"]), "g.csv")
    assert "trop volumineux" in str(exc.value)


def test_trop_de_lignes_refuse(tmp_path, monkeypatch):
    monkeypatch.setattr(classeur, "LIMITE_LIGNES", 5)
    (tmp_path / "g.csv").write_text("a\n" + "\n".join(str(i) for i in range(6)), encoding="utf-8")
    with pytest.raises(ErreurClasseur) as exc:
        charger(faire_ctx(tmp_path, ["g.csv"]), "g.csv")
    assert "5" in str(exc.value) and "lignes" in str(exc.value)
    ecrire_xlsx(tmp_path / "g.xlsx", [["a"]] + [[i] for i in range(6)])
    with pytest.raises(ErreurClasseur):
        charger(faire_ctx(tmp_path, ["g.xlsx"]), "g.xlsx")


def test_fichier_corrompu_message_lisible(tmp_path):
    (tmp_path / "k.xlsx").write_bytes(b"ce n'est pas un classeur")
    with pytest.raises(ErreurClasseur) as exc:
        charger(faire_ctx(tmp_path, ["k.xlsx"]), "k.xlsx")
    assert "illisible" in str(exc.value)


def test_xlsm_lu_sans_macros(tmp_path, monkeypatch):
    vu = {}
    reel = openpyxl.load_workbook

    def espion(*a, **k):
        vu.update(k)
        return reel(*a, **k)
    monkeypatch.setattr(classeur.openpyxl, "load_workbook", espion)
    ecrire_xlsx(tmp_path / "m.xlsx", [["A"], [1]])
    os.rename(tmp_path / "m.xlsx", tmp_path / "m.xlsm")
    # openpyxl refuse un .xlsm sans vbaProject sous certaines versions : le test vérifie l'appel, pas la lecture
    try:
        charger(faire_ctx(tmp_path, ["m.xlsm"]), "m.xlsm")
    except ErreurClasseur:
        pass
    assert vu.get("keep_vba") is False and vu.get("data_only") is True and vu.get("read_only") is True
