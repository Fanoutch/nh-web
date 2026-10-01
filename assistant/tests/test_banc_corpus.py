from docx import Document

from banc.fabriquer_corpus import ecrire_docx, ecrire_txt, lire_source

SOURCE = """# 1. OBJET
Premier paragraphe.

# 1.1. Domaine
Avant la dépose, l'opérateur doit [SAUT] purger le circuit.
[SAUT]
| Tâche | Échéance |
| Visite V1 | 25 h |
"""


def test_lire_source_reconnait_tous_les_blocs():
    blocs = lire_source(SOURCE)
    assert blocs[0] == ("titre", "1. OBJET")
    assert blocs[1] == ("paragraphe", ["Premier paragraphe."])
    assert blocs[3] == ("paragraphe", ["Avant la dépose, l'opérateur doit", None, "purger le circuit."])
    assert blocs[4] == ("saut", None)
    assert blocs[5] == ("tableau", [["Tâche", "Échéance"], ["Visite V1", "25 h"]])


def test_ecrire_docx_contient_titres_saut_et_tableau(tmp_path):
    chemin = ecrire_docx(lire_source(SOURCE), tmp_path / "d.docx")
    doc = Document(str(chemin))
    textes = [p.text for p in doc.paragraphs]
    assert "1.1. Domaine" in textes
    assert doc.tables[0].cell(1, 1).text == "25 h"
    assert 'w:br w:type="page"' in doc.element.xml
    assert "DOCUMENT FICTIF — BANC DE TEST" in doc.sections[0].header.paragraphs[0].text


def test_ecrire_txt_sans_saut_de_page(tmp_path):
    texte = ecrire_txt(lire_source(SOURCE), tmp_path / "d.txt").read_text(encoding="utf-8")
    assert texte.startswith("DOCUMENT FICTIF — BANC DE TEST")
    assert "\f" not in texte
    assert "l'opérateur doit purger le circuit." in texte
    assert "Visite V1 | 25 h" in texte
