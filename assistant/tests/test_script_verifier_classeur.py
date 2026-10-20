import importlib.util
from pathlib import Path

import openpyxl
import pytest

CHEMIN = Path(__file__).resolve().parents[2] / "scripts" / "verifier-classeur.py"
spec = importlib.util.spec_from_file_location("verifier_classeur", CHEMIN)
verifier = importlib.util.module_from_spec(spec)
spec.loader.exec_module(verifier)


def decrire(chemin):
    return "\n".join(verifier.decrire_classeur(chemin))


def test_classeur_propre_ne_montre_aucune_valeur(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stock"
    ws.append(["Réf", "Stock"])
    ws.append(["SECRET-123", 42])
    chemin = tmp_path / "a.xlsx"
    wb.save(chemin)

    texte = decrire(chemin)

    assert "Feuille « Stock » : 2 ligne(s) × 2 colonne(s)" in texte and "colonnes : Réf, Stock" in texte
    assert "SECRET-123" not in texte and "42" not in texte.replace("2 ligne(s) × 2", "")
    assert "⚠" not in texte


def test_entete_pas_en_premiere_ligne_cellules_fusionnees_et_deux_tableaux(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Titre du rapport"])
    ws.merge_cells("A1:C1")
    ws.append([])
    ws.append(["Réf", None, "Qté"])
    ws.append(["A", None, 1])
    ws.append([])
    ws.append(["Réf2", "Qté2"])
    ws.append(["B", 2])
    chemin = tmp_path / "b.xlsx"
    wb.save(chemin)

    texte = decrire(chemin)

    assert "colonnes : Titre du rapport" in texte      # c'est ce que les outils prendraient pour l'en-tête : à signaler
    assert "plage(s) de cellules fusionnées" in texte
    assert "plusieurs blocs" in texte


def test_en_tete_decale_est_signale(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([])
    ws.append([])
    ws.append(["Réf", "Qté"])
    ws.append(["A", 1])
    chemin = tmp_path / "c.xlsx"
    wb.save(chemin)

    assert "en-tête probable en ligne 3 (pas la première)" in decrire(chemin)


def test_formule_sans_resultat_enregistre_est_signalee(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Total"])
    ws.append(["=1+1"])          # écrit par openpyxl : aucun résultat calculé
    chemin = tmp_path / "d.xlsx"
    wb.save(chemin)

    texte = decrire(chemin)

    assert "1 formule(s), dont 1 sans résultat enregistré" in texte and "⚠" in texte


def test_colonne_sans_nom_et_feuille_vide(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(["Réf", None])
    wb.active.append(["A", 1])
    wb.create_sheet("Vide")
    chemin = tmp_path / "e.xlsx"
    wb.save(chemin)

    texte = decrire(chemin)

    assert "au moins une colonne sans nom" in texte and "feuille vide" in texte


def test_fichier_illisible(tmp_path):
    chemin = tmp_path / "f.xlsx"
    chemin.write_bytes(b"pas un classeur")

    assert "illisible" in decrire(chemin)


def test_csv_sans_valeurs_avec_separateur_et_encodage(tmp_path):
    chemin = tmp_path / "g.csv"
    chemin.write_text("Réf;Qté\nCONFIDENTIEL;5\n", encoding="cp1252")

    texte = "\n".join(verifier.decrire_csv(chemin))

    assert "séparateur « ; »" in texte and "colonnes : Réf, Qté" in texte and "CONFIDENTIEL" not in texte


def test_main_signale_un_fichier_introuvable(tmp_path, capsys):
    assert verifier.main([str(tmp_path / "absent.xlsx")]) == 1
    assert "introuvable" in capsys.readouterr().out
    assert verifier.main([]) == 2
