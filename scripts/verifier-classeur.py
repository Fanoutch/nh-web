#!/usr/bin/env python3
"""Décrit la STRUCTURE d'un ou plusieurs classeurs (Excel, CSV), sans jamais afficher une valeur de cellule.

Sert à vérifier sur de VRAIS fichiers du bureau ce que les essais sur fichiers fabriqués ne montrent pas :
ligne d'en-tête qui n'est pas la première, cellules fusionnées, plusieurs tableaux sur une feuille, formules
dont le résultat n'a jamais été calculé (lues comme vides), colonnes sans nom.

Usage (depuis la racine du projet, avec l'environnement de l'assistant, qui a openpyxl) :
    assistant/.venv/bin/python scripts/verifier-classeur.py fichier1.xlsx [fichier2.xlsx ...]

Ne modifie rien, n'écrit rien, n'envoie rien. Les noms de colonnes (en-têtes) sont affichés ; les données, jamais.
"""
from __future__ import annotations

import csv
import sys
from pathlib import Path

MAX_LIGNES = 100_000
MAX_COLONNES_AFFICHEES = 30


def _nom(valeur, maxi=40) -> str:
    texte = " ".join(str(valeur).split())
    return texte if len(texte) <= maxi else texte[:maxi] + "…"


def _vide(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _blocs(lignes_non_vides: list[bool]) -> int:
    """Nombre de blocs de lignes remplies séparés par au moins une ligne vide (≥ 2 : plusieurs tableaux ?)."""
    blocs, dans = 0, False
    for plein in lignes_non_vides:
        if plein and not dans:
            blocs += 1
        dans = plein
    return blocs


def decrire_csv(chemin: Path) -> list[str]:
    for encodage in ("utf-8-sig", "cp1252"):
        try:
            texte = chemin.read_text(encoding=encodage)
            break
        except UnicodeDecodeError:
            continue
    else:
        return ["  ✕ encodage illisible"]
    try:
        dialecte = csv.Sniffer().sniff(texte[:4096], delimiters=";,\t|")
        sep = dialecte.delimiter
    except csv.Error:
        sep = ";"
    lignes = list(csv.reader(texte.splitlines(), delimiter=sep))
    pleines = [any(c.strip() for c in l) for l in lignes]
    premiere = pleines.index(True) if True in pleines else None
    sortie = [f"  encodage {encodage}, séparateur « {sep} », {len(lignes)} ligne(s), {_blocs(pleines)} bloc(s) de lignes"]
    if premiere is None:
        return sortie + ["  ⚠ fichier vide"]
    entete = lignes[premiere]
    sortie.append(f"  en-tête probable : ligne {premiere + 1} ({len(entete)} colonne(s))" + ("  ⚠ pas la première ligne" if premiere else ""))
    sortie.append("  colonnes : " + ", ".join(_nom(c) or "(sans nom)" for c in entete[:MAX_COLONNES_AFFICHEES]))
    if any(not c.strip() for c in entete):
        sortie.append("  ⚠ au moins une colonne sans nom")
    return sortie


def decrire_classeur(chemin: Path) -> list[str]:
    import openpyxl

    try:
        formules = openpyxl.load_workbook(chemin, data_only=False, keep_vba=False)
        valeurs = openpyxl.load_workbook(chemin, data_only=True, keep_vba=False)
    except Exception as exc:  # noqa: BLE001
        return [f"  ✕ illisible ({type(exc).__name__}) : fichier protégé par mot de passe ou abîmé ?"]
    sortie = [f"  {len(formules.sheetnames)} feuille(s)"]
    for nom in formules.sheetnames:
        wf, wv = formules[nom], valeurs[nom]
        lignes_f = list(wf.iter_rows(min_row=1, max_row=min(wf.max_row, MAX_LIGNES), values_only=True))
        lignes_v = list(wv.iter_rows(min_row=1, max_row=min(wv.max_row, MAX_LIGNES), values_only=True))
        pleines = [any(not _vide(c) for c in l) for l in lignes_f]
        entete_idx = pleines.index(True) if True in pleines else None
        sortie.append(f"  Feuille « {_nom(nom)} » : {wf.max_row} ligne(s) × {wf.max_column} colonne(s)"
                      + (" ⚠ tronquée à 100 000 lignes" if wf.max_row > MAX_LIGNES else "") + f", {_blocs(pleines)} bloc(s)")
        if entete_idx is None:
            sortie.append("    ⚠ feuille vide")
            continue
        entete = lignes_f[entete_idx]
        if entete_idx:
            sortie.append(f"    ⚠ en-tête probable en ligne {entete_idx + 1} (pas la première) : les lignes avant seront ignorées")
        sortie.append("    colonnes : " + ", ".join(_nom(c) if not _vide(c) else "(sans nom)" for c in entete[:MAX_COLONNES_AFFICHEES]))
        if any(_vide(c) for c in entete[: wf.max_column]):
            sortie.append("    ⚠ au moins une colonne sans nom")
        if _blocs(pleines) > 1:
            sortie.append("    ⚠ plusieurs blocs séparés par des lignes vides : plusieurs tableaux sur la feuille ?")
        fusionnees = len(wf.merged_cells.ranges)
        if fusionnees:
            sortie.append(f"    ⚠ {fusionnees} plage(s) de cellules fusionnées (seule la première cellule contient la valeur)")
        nb_formules = sans_valeur = 0
        for lf, lv in zip(lignes_f, lignes_v):
            for cf, cv in zip(lf, lv):
                if isinstance(cf, str) and cf.startswith("="):
                    nb_formules += 1
                    sans_valeur += cv is None
        if nb_formules:
            sortie.append(f"    {nb_formules} formule(s)" + (f", dont {sans_valeur} sans résultat enregistré (lues comme vides : ouvrir et enregistrer le fichier dans Excel)" if sans_valeur else ", résultats enregistrés"))
            if sans_valeur:
                sortie[-1] = sortie[-1].replace("    ", "    ⚠ ", 1)
    return sortie


def main(argv: list[str]) -> int:
    if not argv:
        print(__doc__)
        return 2
    code = 0
    for arg in argv:
        chemin = Path(arg)
        print(f"\n=== {chemin.name} ({chemin.stat().st_size / 1024:.0f} Ko) ===" if chemin.is_file() else f"\n=== {arg} : introuvable ===")
        if not chemin.is_file():
            code = 1
            continue
        if chemin.suffix.lower() == ".csv":
            print("\n".join(decrire_csv(chemin)))
        elif chemin.suffix.lower() in (".xlsx", ".xlsm"):
            print("\n".join(decrire_classeur(chemin)))
        else:
            print("  format non pris en charge (attendu : .xlsx, .xlsm, .csv)")
            code = 1
    return code


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
